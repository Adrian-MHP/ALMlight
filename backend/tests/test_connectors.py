"""Connectors: Enterprise Architect (.qea and XMI import with mapping, identities, conflicts, package headings, XMI
export round trip, write-back into .qea), medini analyze (HARA with ISO 26262 ASIL, FMEA with German headers,
Excel export round trip), SysML v2 API (settings, import with satisfy/part-of), permissions and catalogue."""

import io
import json
import sqlite3
import tempfile

import connector_fixtures as fx
from conftest import new_item

P = "/projects/{}"


def _ea(api, project, data, name="model.qea", expect=200, **form):
    files = {"file": (name, data, "application/octet-stream")}
    body = {k: (json.dumps(v) if isinstance(v, dict) else str(v).lower() if isinstance(v, bool) else v)
            for k, v in form.items()}
    return api.post(f"/projects/{project}/connectors/ea/import", files=files, data=body, json=None, expect=expect)


def _items(api, project, **params):
    qs = "&".join(f"{k}={v}" for k, v in params.items())
    return {i["title"]: i for i in api.get(f"/projects/{project}/items?limit=200&{qs}")["items"]}


def test_enterprise_architect_import_mapping_identity_and_conflicts(team, project):
    c, v = team["contributor"], team["viewer"]
    assert {x["key"] for x in c.get("/connectors")} >= {"git", "ea", "medini", "sysml", "reqif"}
    _ea(v, project, fx.ea_qea(), expect=403)
    spec = new_item(c, project, "specification", "EA model")["key"]
    preview = _ea(c, project, fx.ea_qea())
    kinds = {t["kind"]: t for t in preview["analysis"]["types"]}
    assert preview["dry_run"] and not preview["applied"]
    assert kinds["Requirement«functionalRequirement»"]["mapped"] == "system_requirement"
    assert kinds["Requirement"]["count"] == 1 and kinds["Class«block»"]["mapped"] == "system_element"
    assert kinds["Class«testCase»"]["mapped"] == "test_case" and kinds["Activity"]["mapped"] == "function"
    assert "Note" not in kinds and "Package" not in kinds  # diagram notes and package rows are no elements
    assert kinds["Requirement«functionalRequirement»"]["attributes"]["Priority"] == "priority"
    rels = {r["kind"]: r for r in preview["analysis"]["relations"]}
    assert rels["Abstraction«satisfy»"]["mapped"] == "satisfies" and rels["Composition"]["mapped"] == "contains:reverse"
    assert rels["Association"]["mapped"] is None
    assert {p["path"] for p in preview["analysis"]["packages"]} >= {"Requirements / System", "Architecture"}
    assert c.get(f"/projects/{project}/items?q=gauge")["total"] == 0  # nothing saved by the preview
    mapping = {"types": {"Requirement": "stakeholder_requirement"}, "document": spec}
    result = _ea(c, project, fx.ea_qea(), dry_run=False, mapping=mapping)
    assert result["applied"] and result["summary"]["created"] == 8 and result["summary"]["links"] == 6, result
    items = _items(c, project)
    sys1 = items["Measure gauge pressure"]
    assert sys1["type_key"] == "system_requirement" and items["Monitor process pressure"]["type_key"] == "stakeholder_requirement"
    detail = c.get(f"/projects/{project}/items/{sys1['key']}")
    assert detail["fields"]["priority"] == "must" and "gauge pressure" in detail["description"]
    assert "<strong>gauge pressure</strong>" in detail["description"]
    refs = c.get(f"/projects/{project}/items/{sys1['key']}/external-refs")
    assert refs[0]["system"] == "ea" and refs[0]["external_id"] == fx.G["sys1"]
    links = {(lk["type"], lk["direction"], lk["other"]["title"]) for lk in c.get(f"/projects/{project}/items/{sys1['key']}/links")}
    assert ("satisfies", "incoming", "Sensor module") in links and ("verifies", "incoming", "Full-scale test") in links
    assert ("derives_from", "outgoing", "Monitor process pressure") in links
    block_links = {(lk["type"], lk["direction"], lk["other"]["title"]) for lk in
                   c.get(f"/projects/{project}/items/{items['Sensor module']['key']}/links")}
    assert ("contains", "outgoing", "Measuring cell") in block_links and ("realizes", "outgoing", "Measure pressure") in block_links
    # new items sit in the chosen document under a heading per EA package
    tree = c.get(f"/projects/{project}/items/{spec}/document")["root"]
    assert [h["title"] for h in tree["children"]] == ["Requirements"]  # requirements only; architecture has its own tree
    assert [h["title"] for h in tree["children"][0]["children"]] == ["Stakeholder", "System"]
    # re-import: unchanged; a changed note updates the item with a new version
    again = _ea(c, project, fx.ea_qea(), dry_run=False, mapping=mapping)
    assert again["summary"]["unchanged"] == 8 and again["summary"]["created"] == 0 and again["summary"]["links"] == 0
    changed = _ea(c, project, fx.ea_qea("The transmitter shall measure gauge pressure from 0 to 16 bar."),
                  dry_run=False, mapping=mapping)
    assert changed["summary"]["updated"] == 1
    assert "16 bar" in c.get(f"/projects/{project}/items/{sys1['key']}")["description"]
    # edited in Ariadne and in EA: the source is not applied, the conflict is reported
    c.patch(f"/projects/{project}/items/{sys1['key']}", {"title": "Measure gauge pressure (Ariadne)"})
    conflict = _ea(c, project, fx.ea_qea("Range 0 to 25 bar."), dry_run=False, mapping=mapping)
    assert conflict["summary"]["conflicts"] == 1 and "changed in Ariadne and in the source" in conflict["warnings"][0]
    assert c.get(f"/projects/{project}/items/{sys1['key']}")["title"] == "Measure gauge pressure (Ariadne)"
    # wrong files are explained
    err = _ea(c, project, b"not a model", name="x.eap", expect=422)
    assert ".qea" in err["error"]["message"]
    _ea(c, project, b"PK\x03\x04", name="x.txt", expect=422)


def test_enterprise_architect_xmi_import_export_and_writeback(team, project):
    c = team["contributor"]
    r = _ea(c, project, fx.ea_xmi(), name="export.xmi", dry_run=False)
    assert r["applied"] and r["summary"]["created"] == 4, r
    items = _items(c, project)
    req = c.get(f"/projects/{project}/items/{items['Measure gauge pressure']['key']}")
    assert req["type_key"] == "system_requirement" and "shall measure gauge pressure" in req["description"]
    assert c.get(f"/projects/{project}/items/{items['Report diagnostics']['key']}")["description"].count("NAMUR") == 1
    assert r["summary"]["links"] == 3  # satisfy and verify from the EA extension, the plain UML dependency
    # XMI export: EA structure, identities kept, a new Ariadne requirement gets a deterministic EA GUID
    new = new_item(c, project, "system_requirement", "Withstand vibration", fields={"priority": "should"})["key"]
    c.post(f"/projects/{project}/links", {"type": "satisfies", "source": items["Sensor module"]["key"], "target": new},
           expect=201)
    xmi = c.get(f"/projects/{project}/connectors/ea/export.xmi").content
    assert b'extender="Enterprise Architect"' in xmi and fx._xid(fx.G["sys1"]).encode() in xmi
    assert b'stereotype="satisfy"' in xmi and f'value="{new}"'.encode() in xmi
    from almlight.modules.connectors import ea

    model = ea.parse_xmi("roundtrip.xmi", xmi)
    names = {o.name for o in model.objects}
    assert {"Measure gauge pressure", "Withstand vibration", "Sensor module"} <= names
    assert any(r.kind == "Abstraction«satisfy»" for r in model.relations)
    again = _ea(c, project, xmi, name="roundtrip.xmi", dry_run=False)
    assert again["summary"]["created"] == 0, again  # everything maps back to its item
    # write-back into a .qea: imported elements are updated, the new requirement is created with its key
    qea = fx.ea_qea()
    resp = c.post(f"/projects/{project}/connectors/ea/writeback", files={"file": ("model.qea", qea, "x")},
                  data={"package": fx.G["pkg_sys"]}, json=None, expect=200)
    counts = json.loads(resp.headers["x-ariadne-counts"])
    assert counts["created"] >= 1 and counts["connectors"] >= 1
    with tempfile.NamedTemporaryFile(suffix=".qea") as tmp:
        tmp.write(resp.content)
        tmp.flush()
        con = sqlite3.connect(tmp.name)
        row = con.execute("SELECT Name, Alias, Package_ID, Object_Type FROM t_object WHERE Alias = ?", (new,)).fetchone()
        assert row == ("Withstand vibration", new, 4, "Requirement")
        assert con.execute("SELECT Value FROM t_objectproperties WHERE Property='Ariadne key' AND Value=?", (new,)).fetchone()
        assert con.execute("SELECT COUNT(*) FROM t_object").fetchone()[0] > 14
        con.close()
    back = _ea(c, project, resp.content, name="model-ariadne.qea", dry_run=False)
    assert all(t != "Withstand vibration" for t in [o for o in back["warnings"]])
    assert len([i for i in _items(c, project).values() if i["title"] == "Withstand vibration"]) == 1
    bad = c.post(f"/projects/{project}/connectors/ea/writeback", files={"file": ("model.qea", qea, "x")},
                 data={"package": "{00000000-0000-0000-0000-000000000000}"}, json=None, expect=422)
    assert "package" in bad["error"]["message"]


def _sqlite(script: str, base: bytes | None = None) -> bytes:
    with tempfile.NamedTemporaryFile(suffix=".qea") as tmp:
        if base:
            tmp.write(base)
            tmp.flush()
        con = sqlite3.connect(tmp.name)
        con.executescript(script)
        con.close()
        return open(tmp.name, "rb").read()


def test_enterprise_architect_rejects_hostile_model_files(team, project, monkeypatch):
    """Uploaded model files are untrusted: other SQLite files, endless views, odd column names, entity bombs."""
    from almlight.modules.connectors import ea

    c = team["contributor"]
    wb = f"/projects/{project}/connectors/ea/writeback"
    other = _sqlite("CREATE TABLE notes (x TEXT);")
    assert "t_object" in _ea(c, project, other, expect=422)["error"]["message"]
    err = c.post(wb, files={"file": ("x.qea", other, "x")}, data={"package": fx.G["pkg_sys"]}, json=None, expect=422)
    assert "t_object" in err["error"]["message"]
    endless = ("CREATE TABLE t_object (Object_ID INTEGER, Object_Type TEXT, Name TEXT, Note TEXT, Package_ID INTEGER, "
               "Stereotype TEXT, ea_guid TEXT); CREATE VIEW t_package AS WITH RECURSIVE n(x) AS (SELECT 1 UNION ALL "
               "SELECT x + 1 FROM n) SELECT x AS Package_ID, 'P' AS Name, 0 AS Parent_ID, 'g' AS ea_guid FROM n")
    monkeypatch.setattr(ea, "SQLITE_MAX_ROWS", 1000)
    assert "too large" in _ea(c, project, _sqlite(endless + ";"), expect=422)["error"]["message"]
    monkeypatch.setattr(ea, "SQLITE_SECONDS", 0.2)
    assert "too long" in _ea(c, project, _sqlite(endless + " WHERE x < 0;"), expect=422)["error"]["message"]
    monkeypatch.undo()
    # a column name with a quote is written back as a column, not as SQL
    odd = _sqlite('ALTER TABLE t_object ADD COLUMN "Odd"", x" TEXT;', base=fx.ea_qea())
    new_item(c, project, "system_requirement", "Survive salt spray")
    resp = c.post(wb, files={"file": ("odd.qea", odd, "x")}, data={"package": fx.G["pkg_sys"]}, json=None)
    assert json.loads(resp.headers["x-ariadne-counts"])["created"] >= 1
    bomb = (b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">'
            b'<!ENTITY c "&b;&b;&b;&b;&b;&b;&b;&b;&b;&b;"><!ENTITY d "&c;&c;&c;&c;&c;&c;&c;&c;&c;&c;">]>'
            b'<xmi:XMI xmlns:xmi="http://schema.omg.org/spec/XMI/2.1"><e name="&d;&d;&d;&d;&d;"/></xmi:XMI>')
    r = _ea(c, project, bomb, name="bomb.xmi", expect=(200, 422))
    assert "aaaaaaaaaa" not in json.dumps(r)


def _medini(api, project, kind, name, data, expect=200, **form):
    body = {"kind": kind, **{k: (str(v).lower() if isinstance(v, bool) else json.dumps(v) if isinstance(v, dict) else v)
                             for k, v in form.items()}}
    return api.post(f"/projects/{project}/connectors/medini/import", files={"file": (name, data, "x")}, data=body,
                    json=None, expect=expect)


def test_medini_hara_and_fmea(team, project):
    c, v = team["contributor"], team["viewer"]
    preview = _medini(c, project, "hara", "hara.xlsx", fx.hara_xlsx())
    kinds = {t["kind"]: t["mapped"] for t in preview["analysis"]["types"]}
    assert kinds == {"Hazardous event": "hazard", "Safety goal": "system_requirement"}
    r = _medini(c, project, "hara", "hara.xlsx", fx.hara_xlsx(), dry_run=False)
    assert r["summary"]["created"] == 4 and r["summary"]["links"] == 2, r
    haz = {i["title"]: i for i in c.get(f"/projects/{project}/items?type=hazard")["items"]}
    he1 = c.get(f"/projects/{project}/items/{haz['HE-001 Overpressure not detected, vessel bursts']['key']}")
    assert (he1["fields"]["s"], he1["fields"]["e"], he1["fields"]["c"], he1["fields"]["asil"]) == (3, 4, 3, "d")
    he2 = c.get(f"/projects/{project}/items/{haz['HE-002 Overfill not detected']['key']}")
    assert he2["fields"]["asil"] == "b"  # ISO 26262-3 table: S2/E3/C3 = ASIL B
    # HE-003 says QM in the file, but S1/E4/C2 is ASIL A: reported, the computed value counts
    assert any("HE-003" in w and "QM" in w and "gives A" in w for w in r["warnings"]), r["warnings"]
    goal = c.get(f"/projects/{project}/items?type=system_requirement&q=overpressure")["items"][0]
    g = c.get(f"/projects/{project}/items/{goal['key']}")
    assert g["fields"]["asil"] == "d" and g["fields"]["safe_state"] == "Alarm current < 3.6 mA"
    assert g["fields"]["ftti"] == "100 ms" and g["fields"]["safety_relevant"] is True
    assert {lk["other"]["key"] for lk in c.get(f"/projects/{project}/items/{goal['key']}/links")} >= {he1["key"], he2["key"]}
    # FMEA with German headers; the requirement key in the measure becomes a mitigation link
    req = new_item(c, project, "system_requirement", "Compensate temperature drift")["key"]
    f = _medini(c, project, "fmea", "fmea.csv", fx.fmea_csv(req), dry_run=False)
    assert f["summary"]["created"] == 2 and f["summary"]["links"] == 1, f
    risk = c.get(f"/projects/{project}/items?type=risk&q=Drift")["items"][0]
    rd = c.get(f"/projects/{project}/items/{risk['key']}")
    assert (rd["fields"]["severity"], rd["fields"]["occurrence"], rd["fields"]["detection"], rd["fields"]["rpn"]) == (8, 4, 5, 160)
    assert rd["title"] == "Measuring cell: Drift at low temperature"
    # exports in the same layout; re-importing them changes nothing
    hara = v.get(f"/projects/{project}/connectors/medini/hara.xlsx")
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(hara.content)).active
    assert [c_.value for c_ in ws[1]][:8] == ["ID", "Malfunction", "Operational situation", "Hazardous event",
                                               "Severity", "Exposure", "Controllability", "ASIL"]
    assert ws.cell(2, 1).value == "HE-001" and ws.cell(2, 8).value == "D"
    again = _medini(c, project, "hara", "hara-export.xlsx", hara.content, dry_run=False)
    assert again["summary"]["created"] == 0, again
    fmea = v.get(f"/projects/{project}/connectors/medini/fmea.xlsx")
    assert load_workbook(io.BytesIO(fmea.content)).active.cell(2, 3).value == "Drift at low temperature"
    _medini(v, project, "hara", "hara.xlsx", fx.hara_xlsx(), expect=403)
    bad = _medini(c, project, "hara", "x.csv", b"a;b\n1;2\n", expect=422)
    assert "header" in bad["error"]["message"]
    c.post(f"/projects/{project}/connectors/medini/import", files={"file": ("x.csv", b"x", "x")},
           data={"kind": "dfmea"}, json=None, expect=422)


SYSON = {
    "projects": [{"@id": "p1", "@type": "Project", "name": "Drone"}],
    "commits": [{"@id": "c1", "@type": "Commit", "created": "2026-10-01T00:00:00Z"}],
    "elements": [
        {"@id": "pkg", "@type": "Package", "declaredName": "Drone"},
        {"@id": "r1", "@type": "RequirementUsage", "declaredName": "longDistance", "declaredShortName": "REQ-42",
         "reqId": "REQ-42", "text": ["The drone shall have a sufficient range \n * to operate at 5 km."],
         "documentation": [], "qualifiedName": "Drone::Requirements::longDistance", "owner": {"@id": "pkg"}},
        {"@id": "r2", "@type": "RequirementDefinition", "declaredName": "MassRequirement", "text": [],
         "documentation": [{"@id": "d1"}], "qualifiedName": "Drone::MassRequirement", "owner": {"@id": "pkg"}},
        {"@id": "d1", "@type": "Documentation", "body": "The actual mass shall be less than the required mass"},
        {"@id": "p1el", "@type": "PartDefinition", "declaredName": "Drone", "qualifiedName": "Drone::Drone",
         "owner": {"@id": "pkg"}},
        {"@id": "p2el", "@type": "PartUsage", "declaredName": "battery", "qualifiedName": "Drone::Drone::battery",
         "owner": {"@id": "p1el"}},
        {"@id": "anon", "@type": "PartUsage", "declaredName": None, "text": [], "owner": {"@id": "p1el"}},
        {"@id": "s1", "@type": "SatisfyRequirementUsage", "satisfiedRequirement": {"@id": "r1"},
         "satisfyingFeature": {"@id": "p2el"}, "owner": {"@id": "p1el"}},
    ],
}


def test_sysml_v2_connector(team, project, monkeypatch):
    import httpx

    c, m = team["contributor"], team["manager"]

    class Resp:
        def __init__(self, data, status=200):
            self.status_code, self._data = status, data

        def json(self):
            return self._data

    class FakeClient:
        def __init__(self, *a, **k):
            self.headers = k.get("headers", {})

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def get(self, url):
            if url.endswith("/api/rest/projects"):
                return Resp(SYSON["projects"])
            if url.endswith("/commits"):
                return Resp(SYSON["commits"])
            if url.endswith("/elements"):
                return Resp(SYSON["elements"])
            return Resp({}, 404)

    monkeypatch.setattr(httpx, "Client", FakeClient)
    p = f"/projects/{project}"
    c.put(f"{p}/connectors/sysml/settings", {"base_url": "http://syson:8080", "project_id": "p1"}, expect=403)
    m.put(f"{p}/connectors/sysml/settings", {"base_url": "ftp://x"}, expect=422)
    m.put(f"{p}/connectors/sysml/settings", {"base_url": "http://syson:8080", "project_id": "nope"}, expect=422)
    s = m.put(f"{p}/connectors/sysml/settings", {"base_url": "http://syson:8080/", "project_id": "p1", "token": "t0k"})
    assert s["project_name"] == "Drone" and s["has_token"] and s["projects"] == [{"id": "p1", "name": "Drone"}]
    assert "t0k" not in json.dumps(c.get(f"{p}/connectors/sysml/settings"))
    preview = c.post(f"{p}/connectors/sysml/import", {})
    kinds = {t["kind"]: t["mapped"] for t in preview["analysis"]["types"]}
    assert kinds == {"RequirementUsage": "system_requirement", "RequirementDefinition": "system_requirement",
                     "PartDefinition": "system_element", "PartUsage": "system_element"}
    r = c.post(f"{p}/connectors/sysml/import", {"dry_run": False})
    assert r["summary"]["created"] == 4 and r["summary"]["links"] == 2, r
    items = _items(c, project)
    req = c.get(f"{p}/items/{items['REQ-42 longDistance']['key']}")
    assert "sufficient range" in req["description"] and "*" not in req["description"]
    assert "required mass" in c.get(f"{p}/items/{items['MassRequirement']['key']}")["description"]
    battery = items["battery"]["key"]
    links = {(lk["type"], lk["direction"], lk["other"]["title"]) for lk in c.get(f"{p}/items/{battery}/links")}
    assert ("satisfies", "outgoing", "REQ-42 longDistance") in links and ("contains", "incoming", "Drone") in links
    assert c.post(f"{p}/connectors/sysml/import", {"dry_run": False})["summary"]["unchanged"] == 4


def test_upload_imports_preview_when_dry_run_is_omitted(team, project):
    """FastAPI reads a missing boolean form field as false; the import endpoints must still default to a preview."""
    c = team["contributor"]
    csv_data = b"type,title\nsystem_requirement,Preview only\n"
    r = c.post(f"/projects/{project}/import", files={"file": ("x.csv", csv_data, "text/csv")}, data={}, json=None)
    assert r["dry_run"] is True and c.get(f"/projects/{project}/items?q=Preview")["total"] == 0
    r = c.post(f"/projects/{project}/import", files={"file": ("x.csv", csv_data, "text/csv")},
               data={"dry_run": "false"}, json=None)
    assert c.get(f"/projects/{project}/items?q=Preview")["total"] == 1
    m = _medini(c, project, "fmea", "fmea.csv", fx.fmea_csv())
    assert m["dry_run"] is True and not m["applied"]
