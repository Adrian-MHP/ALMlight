"""Ansys medini analyze (functional safety, ISO 26262 / IEC 61508).

Exchange via the formats medini analyze reads and writes:

* **HARA** (hazard analysis and risk assessment) as Excel/CSV: hazardous events with malfunction, operational
  situation, severity S, exposure E, controllability C, ASIL, safety goal, safe state and FTTI → *hazard* items
  (the ASIL is computed from S/E/C as in ISO 26262-3; a different ASIL in the file is reported) and *safety goals*
  (requirements with ASIL, safe state, FTTI) linked by "mitigated by".
* **FMEA** worksheets as Excel/CSV: function/element, failure mode, effect, cause, S/O/D, measures → *risk* items;
  requirement keys mentioned in the measures are linked as mitigations.
* **Safety requirements** via ReqIF (Settings → Import / Export): attributes ASIL, Safe state and FTTI map onto the
  fields of the same names.

Export writes HARA and FMEA in the same column layout, so the tables go back into medini analyze.
"""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from almlight.core.errors import ValidationFailed
from almlight.modules.connectors.base import ExtModel, ExtObject, ExtRelation, Suggestion, norm
from almlight.modules.integrations.models import ExternalRef
from almlight.modules.items.models import Item
from almlight.modules.links.models import Link
from almlight.modules.projects.models import Project

SYSTEM = "medini"
KEY_RE = re.compile(r"(?<![A-Za-z0-9])([A-Z][A-Z0-9]{0,7}-\d+)(?![0-9])")

HARA = {  # canonical column -> accepted headers (English / German, as written by medini analyze and in practice)
    "id": ["id", "hazardous event id", "he id", "event id", "hazard id", "no", "nr", "number"],
    "malfunction": ["malfunction", "malfunctioning behavior", "malfunctioning behaviour", "hazard", "fehlfunktion",
                    "gefährdung"],
    "operational_situation": ["operational situation", "situation", "driving situation", "operating mode",
                              "betriebssituation", "fahrsituation"],
    "hazardous_event": ["hazardous event", "hazardous event description", "description", "gefährdungsereignis",
                        "beschreibung"],
    "s": ["severity", "s", "severity s", "schwere", "schweregrad"],
    "e": ["exposure", "e", "exposure e", "häufigkeit", "exposition"],
    "c": ["controllability", "c", "controllability c", "kontrollierbarkeit", "beherrschbarkeit"],
    "asil": ["asil", "asil level", "integrity level"],
    "safety_goal": ["safety goal", "safety goal text", "sicherheitsziel"],
    "safety_goal_id": ["safety goal id", "sg id", "sg", "sicherheitsziel id"],
    "safe_state": ["safe state", "sicherer zustand"],
    "ftti": ["ftti", "fault tolerant time interval", "fehlertolerante zeitspanne"],
}
FMEA = {
    "id": ["id", "no", "nr", "number", "fm id"],
    "element": ["function", "element", "item", "component", "process step", "funktion", "komponente"],
    "failure_mode": ["failure mode", "potential failure mode", "fehlerart", "fehler"],
    "effect": ["failure effect", "effect", "potential effect", "effects of failure", "fehlerfolge", "auswirkung"],
    "severity": ["severity", "s", "sev", "b", "bedeutung"],
    "cause": ["failure cause", "cause", "potential cause", "fehlerursache", "ursache"],
    "occurrence": ["occurrence", "o", "occ", "a", "auftreten"],
    "prevention": ["prevention", "prevention controls", "current prevention controls", "vermeidungsmaßnahme",
                   "vermeidung"],
    "detection_control": ["detection controls", "current detection controls", "detection measure",
                          "entdeckungsmaßnahme"],
    "detection": ["detection", "d", "det", "e", "entdeckung"],
    "rpn": ["rpn", "rpz"],
    "action_priority": ["action priority", "ap", "aufgabenpriorität"],
    "measures": ["recommended action", "recommended actions", "measures", "mitigation", "actions", "maßnahme",
                 "maßnahmen", "empfohlene maßnahme"],
}
ASIL_KEYS = {"qm": "qm", "a": "a", "b": "b", "c": "c", "d": "d"}


def read_table(filename: str, content: bytes, columns: dict[str, list[str]]) -> tuple[list[dict[str, str]], list[str]]:
    """Rows of the first sheet (or CSV) as canonical-column dicts; the header row is the first row in the first
    20 that matches at least three known column names. Returns (rows, unrecognised headers)."""
    if filename.lower().endswith((".xlsx", ".xlsm")) or content[:2] == b"PK":
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ValidationFailed(f"Cannot read the Excel file: {exc}") from None
        ws = wb.worksheets[0]
        raw = [["" if v is None else str(v).strip() for v in row] for row in ws.iter_rows(values_only=True)]
    else:
        text = content.decode("utf-8-sig", "replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        raw = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    lookup = {norm(alias): key for key, aliases in columns.items() for alias in aliases}
    for h, row in enumerate(raw[:20]):
        mapped = [lookup.get(norm(c)) for c in row]
        if sum(1 for m in mapped if m) >= 3:
            header = mapped
            unknown = [c for c, m in zip(row, mapped, strict=False) if c and not m]
            rows = []
            for r in raw[h + 1:]:
                rec = {k: r[i] for i, k in enumerate(header) if k and i < len(r) and r[i] not in ("", None)}
                if rec:
                    rows.append(rec)
            return rows, unknown
    raise ValidationFailed("No header row found — expected columns such as "
                           + ", ".join(a[0] for a in list(columns.values())[1:6]))


def _level(value: str | None, letter: str, top: int) -> int | None:
    m = re.search(rf"(?i)\b{letter}?\s*([0-{top}])\b", str(value or ""))
    return int(m.group(1)) if m else None


def _asil(value: str | None) -> str | None:
    v = norm(value).replace("asil", "")
    return ASIL_KEYS.get(v)


def computed_asil(s: int | None, e: int | None, c: int | None) -> str | None:
    if s is None or e is None or c is None:
        return None
    if 0 in (s, e, c):
        return "qm"
    return {10: "d", 9: "c", 8: "b", 7: "a"}.get(s + e + c, "qm")


def parse_hara(filename: str, content: bytes) -> ExtModel:
    rows, unknown = read_table(filename, content, HARA)
    model = ExtModel(system=SYSTEM, source=filename)
    if unknown:
        model.warnings.append("Columns not taken over: " + ", ".join(unknown[:12]))
    goals: dict[str, ExtObject] = {}
    for n, r in enumerate(rows, start=1):
        hid = r.get("id") or f"HE-{n:03d}"
        s, e, c = _level(r.get("s"), "S", 3), _level(r.get("e"), "E", 4), _level(r.get("c"), "C", 3)
        given, calc = _asil(r.get("asil")), computed_asil(s, e, c)
        if given and calc and given != calc:
            model.warnings.append(f"{hid}: ASIL in the file is {given.upper()}, S{s}/E{e}/C{c} gives {calc.upper()} "
                                  "(ISO 26262-3) — the computed value is used")
        title = r.get("hazardous_event") or r.get("malfunction") or hid
        fields = {"malfunction": r.get("malfunction"), "operational_situation": r.get("operational_situation"),
                  "hazardous_event": r.get("hazardous_event"), "s": s, "e": e, "c": c, "external_id": hid}
        model.objects.append(ExtObject(ext_id=f"hara:{hid}", kind="Hazardous event", name=f"{hid} {title}"[:300],
                                       fields={k: v for k, v in fields.items() if v not in (None, "")}))
        goal_text = r.get("safety_goal")
        if goal_text or r.get("safety_goal_id"):
            gid = r.get("safety_goal_id") or "SG-" + norm(goal_text)[:40]
            goal = goals.get(gid)
            asil = calc or given
            if goal is None:
                goal = ExtObject(ext_id=f"sg:{gid}", kind="Safety goal",
                                 name=(f"{r['safety_goal_id']} " if r.get("safety_goal_id") else "")
                                 + (goal_text or gid),
                                 text=goal_text or "", fields={"safety_relevant": "true", "kind": "safety",
                                                               "safe_state": r.get("safe_state"), "ftti": r.get("ftti"),
                                                               "asil": asil})
                goals[gid] = goal
                model.objects.append(goal)
            elif asil and "abcd".find(asil) > "abcd".find(goal.fields.get("asil") or "q"):
                goal.fields["asil"] = asil  # a safety goal inherits the highest ASIL of its hazardous events
            goal.fields = {k: v for k, v in goal.fields.items() if v not in (None, "")}
            model.relations.append(ExtRelation(kind="mitigated by", source=f"hara:{hid}", target=f"sg:{gid}"))
    return model


def _num(v: str | None, top: int = 10) -> int | None:
    m = re.search(r"\d+", str(v or ""))
    return int(m.group(0)) if m and 1 <= int(m.group(0)) <= top else None


def parse_fmea(filename: str, content: bytes) -> ExtModel:
    rows, unknown = read_table(filename, content, FMEA)
    model = ExtModel(system=SYSTEM, source=filename)
    if unknown:
        model.warnings.append("Columns not taken over: " + ", ".join(unknown[:12]))
    for n, r in enumerate(rows, start=1):
        fid = r.get("id") or f"FM-{n:03d}"
        mode = r.get("failure_mode") or fid
        name = f"{r['element']}: {mode}" if r.get("element") else mode
        measures = r.get("measures") or ""
        mitigation = "\n".join(x for x in (r.get("prevention"), r.get("detection_control"), measures) if x)
        fields = {"failure_mode": mode, "effect": r.get("effect"), "cause": r.get("cause"),
                  "severity": _num(r.get("severity")), "occurrence": _num(r.get("occurrence")),
                  "detection": _num(r.get("detection")), "mitigation": mitigation}
        model.objects.append(ExtObject(ext_id=f"fmea:{fid}", kind="Failure mode", name=name[:300],
                                       fields={k: v for k, v in fields.items() if v not in (None, "")}))
        for key in dict.fromkeys(KEY_RE.findall(measures)):
            model.relations.append(ExtRelation(kind="mitigated by", source=f"fmea:{fid}", target=f"key:{key}"))
    return model


def suggestion(project: Project, model: ExtModel) -> Suggestion:
    s = Suggestion(types={"Hazardous event": "hazard" if project.model.item_type("hazard") else "risk",
                          "Safety goal": "system_requirement", "Failure mode": "risk"},
                   relations={"mitigated by": "mitigated_by"})
    return s


def parse(kind: str, filename: str, content: bytes) -> ExtModel:
    if kind == "hara":
        return parse_hara(filename, content)
    if kind == "fmea":
        return parse_fmea(filename, content)
    raise ValidationFailed("Choose what the file contains: hara or fmea")


# --- export ---------------------------------------------------------------------------------------------------

HARA_HEADERS = ["ID", "Malfunction", "Operational situation", "Hazardous event", "Severity", "Exposure",
                "Controllability", "ASIL", "Safety goal ID", "Safety goal", "Safe state", "FTTI", "Ariadne key"]
FMEA_HEADERS = ["ID", "Function", "Failure mode", "Failure effect", "Severity", "Failure cause", "Occurrence",
                "Detection", "RPN", "Recommended action", "Ariadne key"]


def _xlsx(headers: list[str], rows: list[list[Any]], title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(48, len(h) + 4))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _mitigations(db: Session, project: Project, item_ids: list) -> dict[Any, list[Item]]:
    if not item_ids:
        return {}
    rows = db.execute(select(Link.source_id, Item).join(Item, Item.id == Link.target_id).where(
        Link.source_id.in_(item_ids), Link.type_key == project.model.semantics.mitigation_link,
        Link.deleted_at.is_(None), Item.deleted_at.is_(None))).all()
    out: dict[Any, list[Item]] = {}
    for sid, item in rows:
        out.setdefault(sid, []).append(item)
    return out


def _source_ids(db: Session, project: Project, item_ids: list) -> dict[Any, str]:
    if not item_ids:
        return {}
    return dict(db.execute(select(ExternalRef.item_id, ExternalRef.external_id).where(
        ExternalRef.project_id == project.id, ExternalRef.system == SYSTEM, ExternalRef.kind == "element",
        ExternalRef.item_id.in_(item_ids))).all())


def export_hara(db: Session, project: Project) -> bytes:
    """One row per hazardous event and safety goal; IDs from medini are kept so the table goes back unchanged."""
    hazards = db.execute(select(Item).where(Item.project_id == project.id, Item.type_key == "hazard",
                                            Item.deleted_at.is_(None)).order_by(Item.seq)).scalars().all()
    goals = _mitigations(db, project, [h.id for h in hazards])
    ids = _source_ids(db, project, [h.id for h in hazards] + [g.id for gs in goals.values() for g in gs])
    rows = []
    for h in hazards:
        f = h.fields
        hid = ids.get(h.id, "").removeprefix("hara:") or f.get("external_id") or h.key
        for g in goals.get(h.id) or [None]:  # type: ignore[list-item]
            gid = (ids.get(g.id, "").removeprefix("sg:") or g.key) if g else None
            gtext = g.title.removeprefix(f"{gid} ") if g else None
            rows.append([hid, f.get("malfunction"), f.get("operational_situation"),
                         f.get("hazardous_event") or h.title,
                         f"S{f['s']}" if f.get("s") is not None else None,
                         f"E{f['e']}" if f.get("e") is not None else None,
                         f"C{f['c']}" if f.get("c") is not None else None, (f.get("asil") or "").upper() or None,
                         gid, gtext, g.fields.get("safe_state") if g else None, g.fields.get("ftti") if g else None,
                         h.key])
    return _xlsx(HARA_HEADERS, rows, "HARA")


def export_fmea(db: Session, project: Project) -> bytes:
    risks = db.execute(select(Item).where(Item.project_id == project.id, Item.type_key == "risk",
                                          Item.deleted_at.is_(None)).order_by(Item.seq)).scalars().all()
    mitig = _mitigations(db, project, [r.id for r in risks])
    rows = []
    ids = _source_ids(db, project, [r.id for r in risks])
    for r in risks:
        f = r.fields
        element, _, mode = r.title.partition(": ") if ": " in r.title else ("", "", r.title)
        actions = "; ".join(f"{m.key} {m.title}" for m in mitig.get(r.id, []))
        rows.append([ids.get(r.id, "").removeprefix("fmea:") or r.key, element or None, f.get("failure_mode") or mode,
                     f.get("effect"), f.get("severity"),
                     f.get("cause"), f.get("occurrence"), f.get("detection"), f.get("rpn"),
                     "\n".join(x for x in (f.get("mitigation"), actions) if x) or None, r.key])
    return _xlsx(FMEA_HEADERS, rows, "FMEA")
