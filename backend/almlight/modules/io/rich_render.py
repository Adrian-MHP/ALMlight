"""Rendering rich text (the block model of ``almlight.core.richtext``) to Word, PDF and ReqIF XHTML — and reading
Word content back.

Word mapping (chosen so that a Word round trip is lossless for everything the editor produces):

=====================  =======================================================================
rich text              Word
=====================  =======================================================================
paragraph / heading    paragraph / paragraph style ``ALM Heading n`` (not in the table of contents)
bold … superscript     run properties (b, i, u, strike, vertAlign); code: character style
                       ``ALM Code Char``; highlight: yellow highlight
link                   ``w:hyperlink`` with an external relationship
item reference         run in character style ``ALM Item Ref``
bullet / numbered      ``List Bullet`` / ``List Number`` (levels 2, 3); numbering restarts per list
quote, code block, rule paragraph styles ``ALM Quote``, ``ALM Code``, ``ALM Rule``
image                  inline picture (alt text in the drawing's description) in style ``ALM Image``
table                  ``Table Grid``; header cells shaded; merged cells for col/row spans
=====================  =======================================================================
"""

from __future__ import annotations

import hashlib
import io
from collections.abc import Callable
from html import escape
from typing import Any

from almlight.core.richtext import (
    IMAGE_SRC,
    Atom,
    Block,
    Inline,
    Run,
    _normalize_inlines,
    canonical,
    parse_html,
    to_html,
)
from almlight.modules.items import storage

EMU_PER_PX = 9525
MAX_WIDTH_IN = 6.3
HEADER_FILL = "E8EDF2"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
RT_HYPERLINK = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"


def _w(tag: str) -> str:
    return f"{{{W}}}{tag}"


# --- images ------------------------------------------------------------------------------------------------


def image_bytes(src: str) -> bytes | None:
    m = IMAGE_SRC.match(src or "")
    if not m:
        return None
    path = storage.blob_path(m.group(1))
    return path.read_bytes() if path.exists() else None


def _image_info(data: bytes) -> tuple[int, int, str]:
    from PIL import Image

    with Image.open(io.BytesIO(data)) as im:
        return im.size[0], im.size[1], im.format or ""


def _office_image(data: bytes) -> bytes:
    """Word and PDF engines do not read WebP: convert it to PNG."""
    from PIL import Image

    w, h, fmt = _image_info(data)
    if fmt.upper() in ("PNG", "JPEG", "GIF"):
        return data
    with Image.open(io.BytesIO(data)) as im:
        out = io.BytesIO()
        im.save(out, format="PNG")
        return out.getvalue()


def export_width_px(attrs: dict[str, Any], natural_w: int, max_px: int) -> int:
    return min(int(attrs.get("width") or natural_w or max_px), max_px)


# ============================== Word: writing =================================================================


class DocxWriter:
    def __init__(self, doc: Any, *, max_width_in: float = MAX_WIDTH_IN):
        self.doc = doc
        self.max_width_in = max_width_in
        self._fresh_cells: set[int] = set()
        self._ensure_styles()

    # -- styles --

    def _ensure_styles(self) -> None:
        from docx.enum.style import WD_STYLE_TYPE
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor

        styles = self.doc.styles
        names = {s.name for s in styles}

        def para(name: str, **kw: Any) -> None:
            if name in names:
                return
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = styles["Normal"]
            st.quick_style = True
            f = st.font
            if kw.get("bold"):
                f.bold = True
            if kw.get("italic"):
                f.italic = True
            if kw.get("size"):
                f.size = Pt(kw["size"])
            if kw.get("font"):
                f.name = kw["font"]
            if kw.get("color"):
                f.color.rgb = RGBColor.from_string(kw["color"])
            pf = st.paragraph_format
            if kw.get("indent"):
                pf.left_indent = Inches(kw["indent"])
            pf.space_before = Pt(kw.get("before", 2))
            pf.space_after = Pt(kw.get("after", 4))
            if kw.get("keep_next"):
                pf.keep_with_next = True
            if kw.get("border"):
                ppr = st.element.get_or_add_pPr()
                bdr = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "9CA3AF")):
                    bottom.set(qn(k), v)
                bdr.append(bottom)
                ppr.append(bdr)

        def char(name: str, **kw: Any) -> None:
            if name in names:
                return
            st = styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
            if kw.get("bold"):
                st.font.bold = True
            if kw.get("font"):
                st.font.name = kw["font"]
            if kw.get("color"):
                st.font.color.rgb = RGBColor.from_string(kw["color"])
            if kw.get("underline"):
                st.font.underline = True

        for level, size in zip(range(1, 7), (14, 12.5, 11.5, 11, 10.5, 10), strict=True):
            para(f"ALM Heading {level}", bold=True, size=size, before=8, after=3, keep_next=True)
        para("ALM Code", font="Consolas", size=9, before=2, after=2)
        para("ALM Quote", italic=True, indent=0.35, color="4B5563")
        para("ALM Rule", border=True, before=2, after=6)
        para("ALM Image", before=4, after=4)
        char("ALM Code Char", font="Consolas")
        char("ALM Item Ref", bold=True, color="0E6E76")
        char("ALM Link", color="1C4F9C", underline=True)

    # -- blocks --

    def write(self, container: Any, blocks: list[Block]) -> None:
        for b in blocks:
            self._block(container, b)

    def _para(self, container: Any, style: str | None = None) -> Any:
        cid = id(container)
        if cid in self._fresh_cells:
            self._fresh_cells.discard(cid)
            p = container.paragraphs[0]
            if style:
                p.style = self.doc.styles[style]
            return p
        return container.add_paragraph(style=style)

    def _align(self, p: Any, align: str | None) -> None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        if align:
            p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                           "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}.get(align)

    def _block(self, c: Any, b: Block, list_level: int = 0) -> None:
        k = b.kind
        if k == "p":
            p = self._para(c)
            self._align(p, b.attrs.get("align"))
            self.inlines(p, b.inlines)
        elif k == "h":
            p = self._para(c, f"ALM Heading {min(int(b.attrs.get('level', 2)), 6)}")
            self._align(p, b.attrs.get("align"))
            self.inlines(p, b.inlines)
        elif k in ("ul", "ol"):
            self._list(c, b, list_level)
        elif k == "quote":
            for child in b.children:
                if child.kind == "p":
                    p = self._para(c, "ALM Quote")
                    self.inlines(p, child.inlines)
                else:
                    self._block(c, child)
        elif k == "pre":
            p = self._para(c, "ALM Code")
            for i, line in enumerate(b.text.split("\n")):
                run = p.add_run(line)
                if i:
                    run._r.insert(0, _br())
        elif k == "hr":
            self._para(c, "ALM Rule")
        elif k == "img":
            p = self._para(c, "ALM Image")
            self.picture(p.add_run(), b.attrs, c)
        elif k == "table":
            self._table(c, b)

    def _list(self, c: Any, b: Block, level: int) -> None:
        ordered = b.kind == "ol"
        base = "List Number" if ordered else "List Bullet"
        style = base if level == 0 else f"{base} {min(level + 1, 3)}"
        num_id = self._restart_numbering(style, b.attrs.get("start", 1)) if ordered else None
        ilvl = getattr(self, "_num_levels", {}).get(num_id, 0)

        def item_para() -> Any:
            p = self._para(c, style)
            if num_id is not None:
                _set_num(p, num_id, ilvl)
            return p

        for li in b.children:
            first = True
            for child in li.children:
                if child.kind in ("p", "h") and first:
                    self.inlines(item_para(), child.inlines)
                    first = False
                elif child.kind in ("ul", "ol"):
                    if first:
                        item_para()
                        first = False
                    self._list(c, child, level + 1)
                elif child.kind in ("p", "h"):
                    cont = "List Continue" if level == 0 else f"List Continue {min(level + 1, 3)}"
                    self.inlines(self._para(c, cont), child.inlines)
                else:
                    if first:
                        item_para()
                        first = False
                    self._block(c, child)

    def _restart_numbering(self, style_name: str, start: int) -> int | None:
        """A new numbering instance for this list so each numbered list starts at ``start``."""
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        try:
            style = self.doc.styles[style_name]
            numbering = self.doc.part.numbering_part.element
        except (KeyError, NotImplementedError):
            return None
        ppr = style.element.pPr
        num_pr = ppr.numPr if ppr is not None else None
        if num_pr is None or num_pr.numId is None:
            return None
        base_id = num_pr.numId.val
        abstract = None
        for num in numbering.findall(qn("w:num")):
            if num.get(qn("w:numId")) == str(base_id):
                el = num.find(qn("w:abstractNumId"))
                abstract = el.get(qn("w:val")) if el is not None else None
        if abstract is None:
            return None
        ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        new_id = max(ids) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_id))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), abstract)
        num.append(a)
        ilvl = int(num_pr.ilvl.val) if num_pr.ilvl is not None else 0
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), str(ilvl))
        so = OxmlElement("w:startOverride")
        so.set(qn("w:val"), str(start or 1))
        override.append(so)
        num.append(override)
        numbering.append(num)
        self._num_levels = getattr(self, "_num_levels", {})
        self._num_levels[new_id] = ilvl
        return new_id

    def _table(self, c: Any, b: Block) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        placements = []
        occupied: set[tuple[int, int]] = set()
        ncols = nrows = 0
        for r, tr in enumerate(b.children):
            col = 0
            for cell in tr.children:
                while (r, col) in occupied:
                    col += 1
                cs, rs = int(cell.attrs.get("colspan", 1)), int(cell.attrs.get("rowspan", 1))
                placements.append((r, col, rs, cs, cell))
                for dr in range(rs):
                    for dc in range(cs):
                        occupied.add((r + dr, col + dc))
                col += cs
                ncols = max(ncols, col)
            nrows = max(nrows, r + 1)
        nrows = max([nrows] + [r + rs for r, _, rs, _, _ in placements])
        if not ncols:
            return
        table = c.add_table(rows=nrows, cols=ncols)
        table.style = self.doc.styles["Table Grid"]
        width_in = self.max_width_in / max(ncols, 1)
        first_row_header = bool(b.children) and all(x.attrs.get("header") for x in b.children[0].children)
        if first_row_header:
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            hdr = OxmlElement("w:tblHeader")
            tr_pr.append(hdr)
        for r, col, rs, cs, cell in placements:
            target = table.cell(r, col)
            if rs > 1 or cs > 1:
                target = target.merge(table.cell(r + rs - 1, col + cs - 1))
                for extra in target.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
            if cell.attrs.get("header"):
                tc_pr = target._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), HEADER_FILL)
                tc_pr.append(shd)
            self._fresh_cells.add(id(target))
            inner = DocxWriter.__new__(DocxWriter)
            inner.doc, inner.max_width_in, inner._fresh_cells = self.doc, max(width_in * cs, 1.0), self._fresh_cells
            inner.write(target, cell.children)

    # -- inline --

    def inlines(self, p: Any, inlines: list[Inline]) -> None:
        for it in inlines:
            marks = dict(it.marks)
            if isinstance(it, Atom):
                if it.kind == "br":
                    run = p.add_run()
                    run.add_break()
                else:
                    run = p.add_run(it.key)
                    run.style = self.doc.styles["ALM Item Ref"]
                continue
            if "link" in marks:
                self._hyperlink(p, marks["link"], it.text, marks)
            else:
                self._format(p.add_run(it.text), marks)

    def _format(self, run: Any, marks: dict[str, str]) -> None:
        from docx.enum.text import WD_COLOR_INDEX

        if "strong" in marks:
            run.bold = True
        if "em" in marks:
            run.italic = True
        if "u" in marks:
            run.underline = True
        if "s" in marks:
            run.font.strike = True
        if "sup" in marks:
            run.font.superscript = True
        if "sub" in marks:
            run.font.subscript = True
        if "code" in marks:
            run.style = self.doc.styles["ALM Code Char"]
        if "mark" in marks:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    def _hyperlink(self, p: Any, url: str, text: str, marks: dict[str, str]) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        r_id = p.part.relate_to(url, RT_HYPERLINK, is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), r_id)
        run = p.add_run(text)
        if "code" not in marks:
            run.style = self.doc.styles["ALM Link"]
        self._format(run, {k: v for k, v in marks.items() if k != "link"})
        p._p.remove(run._r)
        h.append(run._r)
        p._p.append(h)

    def picture(self, run: Any, attrs: dict[str, Any], container: Any = None) -> None:
        from docx.shared import Emu

        data = image_bytes(attrs.get("src", ""))
        if data is None:
            run.add_text(f"[image{': ' + attrs['alt'] if attrs.get('alt') else ''}]")
            return
        office = _office_image(data)
        nat_w, nat_h, _ = _image_info(data)
        max_px = int(self.max_width_in * 96)
        w_px = export_width_px(attrs, nat_w, max_px)
        h_px = int(attrs["height"] * w_px / attrs["width"]) if attrs.get("height") and attrs.get("width") else \
            int(nat_h * w_px / nat_w) if nat_w else w_px
        shape = run.add_picture(io.BytesIO(office), width=Emu(w_px * EMU_PER_PX), height=Emu(h_px * EMU_PER_PX))
        doc_pr = shape._inline.docPr
        doc_pr.set("descr", attrs.get("alt") or "")
        if attrs.get("title"):
            doc_pr.set("title", attrs["title"])


def _br() -> Any:
    from docx.oxml import OxmlElement

    return OxmlElement("w:br")


def _set_num(p: Any, num_id: int, ilvl: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    ppr = p._p.get_or_add_pPr()
    old = ppr.find(qn("w:numPr"))
    if old is not None:
        ppr.remove(old)
    num_pr = OxmlElement("w:numPr")
    lvl = OxmlElement("w:ilvl")
    lvl.set(qn("w:val"), str(ilvl))
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    num_pr.extend([lvl, nid])
    ppr.append(num_pr)


def html_to_docx(doc: Any, container: Any, html: str, *, max_width_in: float = MAX_WIDTH_IN) -> None:
    DocxWriter(doc, max_width_in=max_width_in).write(container, parse_html(html))


# ============================== Word: reading =================================================================


ImageBytesResolver = Callable[[bytes, str], str | None]


def sha_image_url(data: bytes, name: str) -> str | None:
    """Resolver for comparisons: the URL an image would get in the store (no database access)."""
    return "/api/v1/images/" + hashlib.sha256(data).hexdigest()


class DocxContext:
    """What the reader needs from the .docx package besides the body: styles, numbering, relationships, media."""

    def __init__(self, style_names: dict[str, str] | None = None, numbering: Any = None,
                 rels: dict[str, tuple[str, str]] | None = None, media: Callable[[str], bytes | None] | None = None):
        self.style_names = style_names or {}
        self.rels = rels or {}  # rId -> (type, target)
        self.media = media or (lambda rid: None)
        self.formats: dict[tuple[str, str], str] = {}
        self.starts: dict[tuple[str, str], int] = {}
        if numbering is not None:
            self._read_numbering(numbering)

    def _read_numbering(self, numbering: Any) -> None:
        abstract: dict[str, dict[str, str]] = {}
        for an in numbering.iter(_w("abstractNum")):
            levels = {}
            for lvl in an.iter(_w("lvl")):
                fmt = lvl.find(_w("numFmt"))
                levels[lvl.get(_w("ilvl"))] = fmt.get(_w("val")) if fmt is not None else "decimal"
            abstract[an.get(_w("abstractNumId"))] = levels
        for num in numbering.iter(_w("num")):
            a = num.find(_w("abstractNumId"))
            if a is None:
                continue
            for ilvl, fmt in abstract.get(a.get(_w("val")), {}).items():
                self.formats[(num.get(_w("numId")), ilvl)] = fmt
            for ov in num.iter(_w("lvlOverride")):
                so = ov.find(_w("startOverride"))
                if so is not None and (so.get(_w("val")) or "").isdigit():
                    self.starts[(num.get(_w("numId")), ov.get(_w("ilvl"), "0"))] = int(so.get(_w("val")))

    @classmethod
    def from_document(cls, doc: Any) -> DocxContext:
        """Context of a python-docx Document (used for in-memory round trips)."""
        style_names = {s.style_id: s.name for s in doc.styles}
        try:
            numbering = doc.part.numbering_part.element
        except NotImplementedError:
            numbering = None
        rels = {rid: (r.reltype, r.target_ref) for rid, r in doc.part.rels.items()}

        def media(rid: str) -> bytes | None:
            rel = doc.part.rels.get(rid)
            return rel.target_part.blob if rel is not None and not rel.is_external else None

        return cls(style_names, numbering, rels, media)


class DocxReader:
    def __init__(self, ctx: DocxContext, resolver: ImageBytesResolver = sha_image_url,
                 max_width_in: float = MAX_WIDTH_IN):
        self.ctx = ctx
        self.resolver = resolver
        self.max_width_in = max_width_in

    def _style(self, p: Any) -> str:
        ppr = p.find(_w("pPr"))
        st = ppr.find(_w("pStyle")) if ppr is not None else None
        sid = st.get(_w("val")) if st is not None else ""
        return self.ctx.style_names.get(sid, sid or "")

    # -- blocks --

    def blocks(self, elements: list[Any], max_width_in: float | None = None) -> list[Block]:
        width = max_width_in or self.max_width_in
        out: list[Block] = []
        list_entries: list[tuple[int, bool, Block]] = []  # (level, ordered, content block)
        list_id: list[str | None] = [None]
        quote: list[Block] = []
        code: list[str] = []

        def flush_list() -> None:
            if list_entries:
                lst = _build_list(list_entries)
                start = self.ctx.starts.get((list_id[0] or "", "0"), 1) if lst.kind == "ol" else 1
                if start != 1:
                    lst.attrs["start"] = start
                out.append(lst)
                list_entries.clear()

        def flush_quote() -> None:
            if quote:
                out.append(Block("quote", children=list(quote)))
                quote.clear()

        def flush_code() -> None:
            if code:
                out.append(Block("pre", text="\n".join(code)))
                code.clear()

        def flush_all(keep: str = "") -> None:
            if keep != "list":
                flush_list()
            if keep != "quote":
                flush_quote()
            if keep != "code":
                flush_code()

        for el in elements:
            tag = el.tag
            if tag == _w("sdt"):
                content = el.find(_w("sdtContent"))
                if content is not None:
                    flush_all()
                    out.extend(self.blocks(list(content), width))
                continue
            if tag == _w("tbl"):
                flush_all()
                out.append(self.table(el, width))
                continue
            if tag != _w("p"):
                continue
            style = self._style(el)
            inlines, images = self.paragraph_content(el, width)
            lst = self._list_info(el, style)
            if style == "ALM Code":
                flush_all("code")
                code.append("".join(it.text if isinstance(it, Run) else "\n" if it.kind == "br" else it.key
                                    for it in inlines))
                continue
            if style == "ALM Rule":
                flush_all()
                out.append(Block("hr"))
                continue
            if style == "ALM Quote":
                flush_all("quote")
                quote.append(Block("p", inlines=inlines))
                continue
            if lst is not None:
                flush_all("list")
                level, ordered, num_id = lst
                if level == 0 and list_entries and (ordered != list_entries[0][1] or (
                        num_id is not None and list_id[0] is not None and num_id != list_id[0])):
                    flush_list()
                if level == 0:
                    list_id[0] = num_id
                list_entries.append((level, ordered, Block("p", inlines=inlines)))
                for img in images:
                    list_entries.append((level, ordered, img))
                continue
            if style.startswith("List Continue") and list_entries:
                level = int(style.rsplit(" ", 1)[1]) - 1 if style[-1].isdigit() else 0
                list_entries.append((level, list_entries[-1][1], Block("p", inlines=inlines, attrs={"cont": True})))
                continue
            flush_all()
            head = _heading_level(style)
            align = self._align(el)
            if inlines or not images:
                attrs: dict[str, Any] = {"align": align} if align else {}
                if head:
                    out.append(Block("h", inlines=inlines, attrs={"level": head, **attrs}))
                else:
                    out.append(Block("p", inlines=inlines, attrs=attrs))
            out.extend(images)
        flush_all()
        return out

    def _align(self, p: Any) -> str | None:
        ppr = p.find(_w("pPr"))
        jc = ppr.find(_w("jc")) if ppr is not None else None
        val = jc.get(_w("val")) if jc is not None else None
        return {"center": "center", "right": "right", "end": "right", "both": "justify",
                "distribute": "justify"}.get(val or "")

    def _list_info(self, p: Any, style: str) -> tuple[int, bool, str | None] | None:
        ppr = p.find(_w("pPr"))
        num_pr = ppr.find(_w("numPr")) if ppr is not None else None
        nid = num_pr.find(_w("numId")) if num_pr is not None else None
        ilvl = num_pr.find(_w("ilvl")) if num_pr is not None else None
        num_id = nid.get(_w("val")) if nid is not None else None
        level = 0
        if style.startswith(("List Bullet", "List Number")):
            ordered = style.startswith("List Number")
            if style[-1].isdigit():
                level = int(style.rsplit(" ", 1)[1]) - 1
            elif ilvl is not None:
                level = int(ilvl.get(_w("val"), "0"))
            return level, ordered, num_id
        if num_id in (None, "0"):
            return None
        lvl = ilvl.get(_w("val"), "0") if ilvl is not None else "0"
        fmt = self.ctx.formats.get((num_id, lvl), "bullet")
        return int(lvl), fmt not in ("bullet", "none"), num_id

    # -- paragraph content --

    def paragraph_content(self, p: Any, width: float) -> tuple[list[Inline], list[Block]]:
        inlines: list[Inline] = []
        images: list[Block] = []
        state = {"field": 0, "skip": False}
        self._content(p, inlines, images, (), state, width)
        return _normalize_inlines(inlines), images

    def _content(self, el: Any, inlines: list[Inline], images: list[Block], marks: tuple, state: dict,
                 width: float) -> None:
        for child in el:
            tag = child.tag
            if tag == _w("r"):
                self._run(child, inlines, images, marks, state, width)
            elif tag == _w("hyperlink"):
                rid = child.get(f"{{{R_NS}}}id")
                rel = self.ctx.rels.get(rid or "")
                from almlight.core.richtext import safe_href

                href = safe_href(rel[1]) if rel and rel[0] == RT_HYPERLINK else None
                self._content(child, inlines, images, marks + ((("link", href),) if href else ()), state, width)
            elif tag in (_w("ins"), _w("smartTag"), _w("fldSimple"), _w("customXml"), _w("dir"), _w("bdo")):
                self._content(child, inlines, images, marks, state, width)
            elif tag == _w("sdt"):
                content = child.find(_w("sdtContent"))
                if content is not None:
                    self._content(content, inlines, images, marks, state, width)

    def _run(self, r: Any, inlines: list[Inline], images: list[Block], marks: tuple, state: dict,
             width: float) -> None:
        from almlight.core.richtext import _order

        rpr = r.find(_w("rPr"))
        run_marks = list(marks)
        itemref = False
        if rpr is not None:
            def on(name: str) -> bool:
                e = rpr.find(_w(name))
                return e is not None and e.get(_w("val"), "true").lower() not in ("0", "false", "none")

            if on("b"):
                run_marks.append(("strong", ""))
            if on("i"):
                run_marks.append(("em", ""))
            u = rpr.find(_w("u"))
            if u is not None and u.get(_w("val"), "single") != "none":
                run_marks.append(("u", ""))
            if on("strike") or on("dstrike"):
                run_marks.append(("s", ""))
            va = rpr.find(_w("vertAlign"))
            if va is not None:
                if va.get(_w("val")) == "superscript":
                    run_marks.append(("sup", ""))
                elif va.get(_w("val")) == "subscript":
                    run_marks.append(("sub", ""))
            hl = rpr.find(_w("highlight"))
            shd = rpr.find(_w("shd"))
            if (hl is not None and hl.get(_w("val")) not in (None, "none")) or (
                    shd is not None and (shd.get(_w("fill")) or "auto").lower() not in ("auto", "ffffff", "none")):
                run_marks.append(("mark", ""))
            rs = rpr.find(_w("rStyle"))
            sname = self.ctx.style_names.get(rs.get(_w("val")), rs.get(_w("val"))) if rs is not None else ""
            fonts = rpr.find(_w("rFonts"))
            font = (fonts.get(_w("ascii")) or "") if fonts is not None else ""
            if sname == "ALM Code Char" or font.lower() in ("consolas", "courier new", "courier"):
                run_marks.append(("code", ""))
            itemref = sname == "ALM Item Ref"
        ordered = _order(run_marks)
        text: list[str] = []
        for c in r:
            tag = c.tag
            if tag == _w("fldChar"):
                kind = c.get(_w("fldCharType"))
                if kind == "begin":
                    state["field"] += 1
                    state["skip"] = True
                elif kind == "separate":
                    state["skip"] = False
                elif kind == "end":
                    state["field"] = max(0, state["field"] - 1)
                    state["skip"] = False
            elif state["skip"] or tag == _w("instrText"):
                continue
            elif tag == _w("t"):
                text.append(c.text or "")
            elif tag == _w("tab"):
                text.append(" ")
            elif tag == _w("noBreakHyphen"):
                text.append("-")
            elif tag in (_w("br"), _w("cr")):
                if c.get(_w("type")) in ("page", "column"):
                    continue
                if text:
                    inlines.append(Run("".join(text), ordered))
                    text = []
                inlines.append(Atom("br", marks=ordered))
            elif tag == _w("drawing"):
                if text:
                    inlines.append(Run("".join(text), ordered))
                    text = []
                img = self._drawing(c, width)
                if img is not None:
                    images.append(img)
        if text:
            joined = "".join(text)
            if itemref and joined.strip():
                import re

                key = joined.strip().upper()
                if re.fullmatch(r"[A-Z][A-Z0-9]{0,7}-\d+", key):
                    inlines.append(Atom("itemref", key=key, marks=ordered))
                    return
            inlines.append(Run(joined, ordered))

    def _drawing(self, d: Any, width: float) -> Block | None:
        blip = next(d.iter(f"{{{A_NS}}}blip"), None)
        if blip is None:
            return None
        rid = blip.get(f"{{{R_NS}}}embed")
        data = self.ctx.media(rid) if rid else None
        if not data:
            return None
        doc_pr = next(d.iter(f"{{{WP_NS}}}docPr"), None)
        name = (doc_pr.get("name") if doc_pr is not None else None) or "image"
        src = self.resolver(data, name)
        if not src:
            return None
        attrs: dict[str, Any] = {"src": src}
        if doc_pr is not None and doc_pr.get("descr"):
            attrs["alt"] = doc_pr.get("descr")[:500]
        if doc_pr is not None and doc_pr.get("title"):
            attrs["title"] = doc_pr.get("title")[:500]
        extent = next(d.iter(f"{{{WP_NS}}}extent"), None)
        try:
            nat_w, nat_h, _ = _image_info(data)
        except Exception:  # noqa: BLE001
            nat_w = nat_h = 0
        if extent is not None and nat_w:
            w_px = round(int(extent.get("cx", "0")) / EMU_PER_PX)
            h_px = round(int(extent.get("cy", "0")) / EMU_PER_PX)
            expected = export_width_px({}, nat_w, int(width * 96))
            if abs(w_px - expected) > 1:
                attrs["width"] = w_px
                if nat_h and abs(h_px - round(nat_h * w_px / nat_w)) > 1:
                    attrs["height"] = h_px
        return Block("img", attrs=attrs)

    # -- tables --

    def table(self, tbl: Any, width: float) -> Block:
        rows: list[Block] = []
        open_cells: dict[int, Block] = {}  # grid column -> cell started with vMerge=restart
        grid_cols = max(1, len(tbl.findall(f"{_w('tblGrid')}/{_w('gridCol')}")))
        for tr in tbl.findall(_w("tr")):
            tr_pr = tr.find(_w("trPr"))
            header_row = tr_pr is not None and tr_pr.find(_w("tblHeader")) is not None
            cells: list[Block] = []
            col = 0
            for tc in tr.findall(_w("tc")):
                tc_pr = tc.find(_w("tcPr"))
                span = 1
                vmerge = None
                fill = ""
                if tc_pr is not None:
                    gs = tc_pr.find(_w("gridSpan"))
                    span = int(gs.get(_w("val"), "1")) if gs is not None else 1
                    vm = tc_pr.find(_w("vMerge"))
                    if vm is not None:
                        vmerge = vm.get(_w("val"), "continue")
                    shd = tc_pr.find(_w("shd"))
                    fill = (shd.get(_w("fill")) or "").upper() if shd is not None else ""
                if vmerge == "continue" and col in open_cells:
                    started = open_cells[col]
                    started.attrs["rowspan"] = started.attrs.get("rowspan", 1) + 1
                    col += span
                    continue
                cell_width = width * span / grid_cols
                content = self.blocks(list(tc), cell_width)
                from almlight.core.richtext import _ensure_blocks, _trim_blocks

                cell = Block("cell", children=_ensure_blocks(_trim_blocks(content)),
                             attrs={"header": header_row or fill == HEADER_FILL})
                if span > 1:
                    cell.attrs["colspan"] = span
                if vmerge == "restart":
                    open_cells[col] = cell
                else:
                    open_cells.pop(col, None)
                cells.append(cell)
                col += span
            if cells:
                rows.append(Block("tr", children=cells))
        return Block("table", children=rows)


def _heading_level(style: str) -> int | None:
    for prefix in ("ALM Heading ", "Heading "):
        if style.startswith(prefix) and style[len(prefix):].isdigit():
            return max(1, min(int(style[len(prefix):]), 6))
    return None


def _build_list(entries: list[tuple[int, bool, Block]]) -> Block:
    """Nested lists from (level, ordered, block) entries in document order."""
    root = Block("ol" if entries[0][1] else "ul")
    stack: list[tuple[int, Block]] = [(0, root)]
    for level, ordered, block in entries:
        level = max(0, level)
        while len(stack) - 1 > level:
            stack.pop()
        while len(stack) - 1 < level:
            parent_list = stack[-1][1]
            if not parent_list.children:
                parent_list.children.append(Block("li", children=[Block("p")]))
            nested = Block("ol" if ordered else "ul")
            parent_list.children[-1].children.append(nested)
            stack.append((len(stack), nested))
        current = stack[-1][1]
        attach = block.attrs.pop("cont", False) or block.kind == "img"
        if attach and current.children:
            current.children[-1].children.append(block)
        else:
            current.children.append(Block("li", children=[block]))
    return root


def docx_blocks_to_html(elements: list[Any], ctx: DocxContext, resolver: ImageBytesResolver = sha_image_url) -> str:
    return canonical(to_html(DocxReader(ctx, resolver).blocks(elements)), lambda s, a: s)


def docx_roundtrip_html(html: str) -> str:
    """What a rich text looks like after export to Word and reading it back (for change detection)."""
    from docx import Document

    doc = Document()
    body = doc.element.body
    before = len(body)
    DocxWriter(doc).write(doc, parse_html(html))
    new = list(body)[before - 1:-1]
    return docx_blocks_to_html(new, DocxContext.from_document(doc))


# ============================== PDF ==================================================================================


def _rl_inline(inlines: list[Inline], mono: str) -> str:
    out: list[str] = []
    for it in inlines:
        marks = dict(it.marks)
        if isinstance(it, Atom):
            out.append("<br/>" if it.kind == "br" else f'<font color="#0e6e76"><b>{escape(it.key)}</b></font>')
            continue
        s = escape(it.text)
        if "code" in marks:
            s = f'<font face="{mono}">{s}</font>'
        if "mark" in marks:
            s = f'<font backColor="#fff3a3">{s}</font>'
        for name, tag in (("sub", "sub"), ("sup", "super"), ("s", "strike"), ("u", "u"), ("em", "i"),
                          ("strong", "b")):
            if name in marks:
                s = f"<{tag}>{s}</{tag}>"
        if "link" in marks:
            s = f'<a href="{escape(marks["link"], quote=True)}" color="#1c4f9c">{s}</a>'
        out.append(s)
    return "".join(out)


def pdf_flowables(html: str, styles: dict[str, Any], *, max_width: float, mono: str = "Courier") -> list[Any]:
    return _pdf_blocks(parse_html(html), styles, max_width, mono)


def _pdf_blocks(blocks: list[Block], s: dict[str, Any], max_width: float, mono: str) -> list[Any]:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import HRFlowable, Image, ListFlowable, ListItem, Paragraph, Preformatted, Table, TableStyle

    out: list[Any] = []
    aligns = {"center": TA_CENTER, "right": TA_RIGHT, "justify": TA_JUSTIFY}

    def style_for(base: Any, align: str | None) -> Any:
        if not align:
            return base
        return ParagraphStyle(f"{base.name}-{align}", parent=base, alignment=aligns[align])

    for b in blocks:
        k = b.kind
        if k == "p":
            out.append(Paragraph(_rl_inline(b.inlines, mono) or "&nbsp;", style_for(s["body"], b.attrs.get("align"))))
        elif k == "h":
            level = int(b.attrs.get("level", 2))
            size = {1: 13, 2: 12, 3: 11, 4: 10.5, 5: 10, 6: 9.5}.get(level, 10)
            hs = ParagraphStyle(f"desc-h{level}", parent=s["body"], fontName=s["req"].fontName, fontSize=size,
                                leading=size + 3, spaceBefore=6, spaceAfter=2)
            out.append(Paragraph(_rl_inline(b.inlines, mono), style_for(hs, b.attrs.get("align"))))
        elif k in ("ul", "ol"):
            items = [ListItem(_pdf_blocks(li.children, s, max_width - 18, mono) or [Paragraph("", s["body"])],
                              leftIndent=14) for li in b.children]
            kw: dict[str, Any] = {"bulletType": "1" if k == "ol" else "bullet", "leftIndent": 14,
                                  "bulletFontSize": 8}
            if k == "ol" and b.attrs.get("start"):
                kw["start"] = b.attrs["start"]
            out.append(ListFlowable(items, **kw))
        elif k == "quote":
            qs = ParagraphStyle("desc-quote", parent=s["body"], leftIndent=14, textColor=colors.HexColor("#4b5563"),
                                fontName=s["body"].fontName)
            for child in b.children:
                if child.kind == "p":
                    out.append(Paragraph(f"<i>{_rl_inline(child.inlines, mono)}</i>", qs))
                else:
                    out.extend(_pdf_blocks([child], s, max_width - 14, mono))
        elif k == "pre":
            out.append(Preformatted(b.text, s["code"]))
        elif k == "hr":
            out.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#9ca3af"), spaceBefore=2,
                                  spaceAfter=4))
        elif k == "img":
            data = image_bytes(b.attrs.get("src", ""))
            if data is None:
                out.append(Paragraph(escape(f"[image: {b.attrs.get('alt', '')}]"), s["meta"]))
                continue
            nat_w, nat_h, _ = _image_info(data)
            w_px = int(b.attrs.get("width") or nat_w)
            h_px = int(b.attrs.get("height") or (nat_h * w_px / nat_w if nat_w else w_px))
            w_pt = w_px * 0.75
            h_pt = h_px * 0.75
            if w_pt > max_width:
                h_pt, w_pt = h_pt * max_width / w_pt, max_width
            img = Image(io.BytesIO(_office_image(data)), width=w_pt, height=h_pt)
            img.hAlign = "LEFT"
            out.append(img)
        elif k == "table":
            out.append(_pdf_table(b, s, max_width, mono, Table, TableStyle, colors))
    return out


def _pdf_table(b: Block, s: dict[str, Any], max_width: float, mono: str, Table: Any, TableStyle: Any,
               colors: Any) -> Any:
    placements = []
    occupied: set[tuple[int, int]] = set()
    ncols = nrows = 0
    for r, tr in enumerate(b.children):
        col = 0
        for cell in tr.children:
            while (r, col) in occupied:
                col += 1
            cs, rs = int(cell.attrs.get("colspan", 1)), int(cell.attrs.get("rowspan", 1))
            placements.append((r, col, rs, cs, cell))
            for dr in range(rs):
                for dc in range(cs):
                    occupied.add((r + dr, col + dc))
            col += cs
            ncols = max(ncols, col)
        nrows = max(nrows, r + 1)
    nrows = max([nrows] + [r + rs for r, _, rs, _, _ in placements])
    col_w = max_width / max(ncols, 1)
    data: list[list[Any]] = [["" for _ in range(ncols)] for _ in range(nrows)]
    style = [("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9ca3af")), ("VALIGN", (0, 0), (-1, -1), "TOP")]
    for r, col, rs, cs, cell in placements:
        data[r][col] = _pdf_blocks(cell.children, {**s, "body": s["cell"]}, col_w * cs - 6, mono) or ""
        if rs > 1 or cs > 1:
            style.append(("SPAN", (col, r), (col + cs - 1, r + rs - 1)))
        if cell.attrs.get("header"):
            style.append(("BACKGROUND", (col, r), (col + cs - 1, r + rs - 1), colors.HexColor("#e8edf2")))
    t = Table(data, colWidths=[col_w] * ncols, hAlign="LEFT", repeatRows=1 if b.children and all(
        c.attrs.get("header") for c in b.children[0].children) else 0)
    t.setStyle(TableStyle(style))
    return t


# ============================== ReqIF XHTML ========================================================================

XHTML = "http://www.w3.org/1999/xhtml"


def xhtml_elements(html: str, add_file: Callable[[str], str | None]) -> list[Any]:
    """Rich text as XHTML elements (ElementTree) for a ReqIF ATTRIBUTE-VALUE-XHTML. ``add_file`` puts an image into
    the archive and returns its path there."""
    from xml.etree import ElementTree as ET

    def x(tag: str, parent: Any = None, **attrs: str) -> Any:
        el = ET.Element(f"{{{XHTML}}}{tag}", attrs) if parent is None else ET.SubElement(parent, f"{{{XHTML}}}{tag}",
                                                                                          attrs)
        return el

    def inline(parent: Any, inlines: list[Inline]) -> None:
        last: Any = None

        def text(t: str) -> None:
            nonlocal last
            if last is None:
                parent.text = (parent.text or "") + t
            else:
                last.tail = (last.tail or "") + t

        for it in inlines:
            if isinstance(it, Atom):
                if it.kind == "br":
                    last = x("br", parent)
                else:
                    last = x("span", parent, title="item reference")
                    last.text = it.key
                continue
            marks = dict(it.marks)
            if not marks:
                text(it.text)
                continue
            outer = None
            cur = None
            for name in ("link", "strong", "em", "u", "s", "sup", "sub", "code", "mark"):
                if name not in marks:
                    continue
                if name == "link":
                    el = ET.Element(f"{{{XHTML}}}a", {"href": marks["link"]})
                elif name in ("u", "s", "mark"):
                    css = {"u": "text-decoration:underline", "s": "text-decoration:line-through",
                           "mark": "background-color:yellow"}[name]
                    el = ET.Element(f"{{{XHTML}}}span", {"style": css})
                else:
                    el = ET.Element(f"{{{XHTML}}}{name}")
                if cur is None:
                    outer = el
                else:
                    cur.append(el)
                cur = el
            cur.text = it.text
            parent.append(outer)
            last = outer

    def blocks(parent: Any, items: list[Block]) -> None:
        for b in items:
            k = b.kind
            if k in ("p", "h"):
                tag = "p" if k == "p" else f"h{min(int(b.attrs.get('level', 2)), 6)}"
                el = x(tag, parent)
                if b.attrs.get("align"):
                    el.set("style", f"text-align:{b.attrs['align']}")
                inline(el, b.inlines)
            elif k in ("ul", "ol"):
                el = x(k, parent)
                for li in b.children:
                    blocks(x("li", el), li.children)
            elif k == "quote":
                blocks(x("blockquote", parent), b.children)
            elif k == "pre":
                x("pre", parent).text = b.text
            elif k == "hr":
                x("hr", parent)
            elif k == "img":
                data = image_bytes(b.attrs.get("src", ""))
                path = add_file(b.attrs["src"]) if data is not None else None
                if path is None:
                    x("p", parent).text = f"[image: {b.attrs.get('alt', '')}]"
                    continue
                from almlight.modules.images.service import EXTENSIONS

                w, h, fmt = _image_info(data)
                ctype = {"PNG": "image/png", "JPEG": "image/jpeg", "GIF": "image/gif", "WEBP": "image/webp"}.get(
                    fmt.upper(), "image/png")
                assert ctype in EXTENSIONS
                el = x("object", parent, data=path, type=ctype)
                if b.attrs.get("width"):
                    el.set("width", str(b.attrs["width"]))
                if b.attrs.get("height"):
                    el.set("height", str(b.attrs["height"]))
                el.text = b.attrs.get("alt") or ""
            elif k == "table":
                t = x("table", parent, border="1")
                for tr in b.children:
                    row = x("tr", t)
                    for cell in tr.children:
                        c = x("th" if cell.attrs.get("header") else "td", row)
                        for n in ("colspan", "rowspan"):
                            if cell.attrs.get(n):
                                c.set(n, str(cell.attrs[n]))
                        blocks(c, cell.children)

    root = ET.Element(f"{{{XHTML}}}div")
    blocks(root, parse_html(html))
    return [root]


def xhtml_to_html(el: Any) -> str:
    """XHTML of a ReqIF value as HTML text; images (``object`` with an image type, or an image fallback nested in
    another object, as DOORS writes them) become ``<img src="<archive path>">`` for the importer to resolve."""
    from xml.etree import ElementTree as ET

    def local(tag: str) -> str:
        return tag.rsplit("}", 1)[-1].lower() if isinstance(tag, str) else ""

    def image_of(obj: Any) -> Any | None:
        typ = (obj.get("type") or "").lower()
        data = obj.get("data") or ""
        if typ.startswith("image/") or data.lower().endswith((".png", ".jpg", ".jpeg", ".gif", ".webp")):
            return obj
        for c in obj:
            if local(c.tag) == "object":
                found = image_of(c)
                if found is not None:
                    return found
        return None

    def rebuild(src: Any) -> Any:
        tag = local(src.tag)
        if tag == "object":
            img = image_of(src)
            if img is None:
                el = ET.Element("span")
                el.text = "".join(src.itertext()).strip()
                el.tail = src.tail
                return el
            attrs = {"src": img.get("data") or ""}
            for n in ("width", "height"):
                if img.get(n):
                    attrs[n] = img.get(n)
            alt = "".join(img.itertext()).strip()
            if alt:
                attrs["alt"] = alt
            el = ET.Element("img", attrs)
            el.tail = src.tail
            return el
        attrs = {k.rsplit("}", 1)[-1]: v for k, v in src.attrib.items()}
        if tag == "span" and attrs.get("title") == "item reference":
            attrs = {"data-type": "item-ref", "data-key": (src.text or "").strip()}
        el = ET.Element(tag or "span", attrs)
        el.text = src.text
        el.tail = src.tail
        for c in src:
            el.append(rebuild(c))
        return el

    parts = [el.text or ""]
    for c in el:
        parts.append(ET.tostring(rebuild(c), encoding="unicode", method="html"))
    return "".join(parts)
