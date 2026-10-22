"""CSV / Excel export and import of items.

Columns: ``key, type, title, state, version, owner, description, applicability, parent`` plus one
column per custom field (header = field key) and one per outgoing link type (header = link type
key, cell = comma-separated target keys). Import accepts the same layout: rows with an existing
``key`` update that item, rows without create one. Everything runs in a single transaction, and
``dry_run`` validates and reports without writing.
"""

from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from almlight.core.errors import DomainError, ValidationFailed
from almlight.modules.identity.models import User
from almlight.modules.items import service as items_svc
from almlight.modules.items.models import Item
from almlight.modules.links import service as links_svc
from almlight.modules.links.models import Link
from almlight.modules.projects.models import Project

BASE_COLUMNS = ["key", "type", "title", "state", "version", "owner", "description", "applicability", "parent"]
MAX_ROWS = 20_000


def _rows_for_export(db: Session, project: Project, types: list[str] | None) -> tuple[list[str], list[list[Any]]]:
    model = project.model
    selected = [t for t in model.item_types if not types or t.key in types]
    stmt = select(Item).where(Item.project_id == project.id, Item.deleted_at.is_(None),
                              Item.type_key.in_([t.key for t in selected])).order_by(Item.type_key, Item.seq)
    items = db.execute(stmt).scalars().all()
    field_keys: list[str] = []
    for t in selected:
        for f in t.fields:
            if f.key not in field_keys:
                field_keys.append(f.key)
    structural = model.structural_link_types()
    link_keys = [lt.key for lt in model.link_types if not lt.structural]
    ids = [i.id for i in items]
    outgoing: dict[Any, dict[str, list[str]]] = {i: {} for i in ids}
    parents: dict[Any, str] = {}
    users = {u.id: u.username for u in db.execute(select(User)).scalars()}
    if ids:
        rows = db.execute(
            select(Link, Item.key).join(Item, Item.id == Link.target_id)
            .where(Link.source_id.in_(ids), Link.deleted_at.is_(None))
        ).all()
        keys = {i.id: i.key for i in items}
        for link, target_key in rows:
            outgoing[link.source_id].setdefault(link.type_key, []).append(target_key)
        for link in db.execute(
            select(Link).where(Link.target_id.in_(ids), Link.type_key.in_(structural), Link.deleted_at.is_(None))
        ).scalars():
            if link.source_id in keys:
                parents.setdefault(link.target_id, keys[link.source_id])
    header = BASE_COLUMNS + field_keys + link_keys
    out = []
    for it in items:
        row = [it.key, it.type_key, it.title, it.state or "", it.version_label, users.get(it.owner_id, ""),
               it.description, it.applicability or "", parents.get(it.id, "")]
        for fk in field_keys:
            v = it.fields.get(fk)
            row.append(", ".join(map(str, v)) if isinstance(v, list) else ("" if v is None else v))
        for lk in link_keys:
            row.append(", ".join(sorted(outgoing[it.id].get(lk, []))))
        out.append(row)
    return header, out


def export_csv(db: Session, project: Project, types: list[str] | None = None) -> bytes:
    header, rows = _rows_for_export(db, project, types)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(header)
    writer.writerows(rows)
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel detects UTF-8


def export_xlsx(db: Session, project: Project, types: list[str] | None = None) -> bytes:
    header, rows = _rows_for_export(db, project, types)
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    ws.append(header)
    for row in rows:
        ws.append([str(v) if isinstance(v, bool) else v for v in row])
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _read_table(filename: str, content: bytes) -> list[dict[str, str]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise ValidationFailed(f"Cannot read Excel file: {exc}") from None
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        try:
            header = [str(h).strip() if h is not None else "" for h in next(it)]
        except StopIteration:
            return []
        rows = []
        for values in it:
            if values is None or all(v is None or str(v).strip() == "" for v in values):
                continue
            rows.append({header[i]: ("" if v is None else str(v)) for i, v in enumerate(values) if i < len(header)})
        return rows
    text = content.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [dict(r) for r in csv.DictReader(io.StringIO(text), dialect=dialect)]


def import_table(
    db: Session, project: Project, actor: User, filename: str, content: bytes, *, default_type: str | None = None,
    dry_run: bool = True,
) -> dict[str, Any]:
    rows = _read_table(filename, content)
    if len(rows) > MAX_ROWS:
        raise ValidationFailed(f"At most {MAX_ROWS} rows per import")
    model = project.model
    link_keys = {lt.key for lt in model.link_types if not lt.structural}
    results: list[dict[str, Any]] = []
    errors = 0
    created: dict[str, Item] = {}
    pending_links: list[tuple[int, Item, str, list[str]]] = []
    pending_parents: list[tuple[int, Item, str]] = []
    savepoint = db.begin_nested()
    try:
        for idx, raw in enumerate(rows, start=2):  # row 1 is the header
            row = {(k or "").strip().lower(): (v or "").strip() for k, v in raw.items() if k}
            try:
                type_key = row.get("type") or default_type
                key = row.get("key", "").upper()
                existing = None
                if key:
                    existing = db.execute(
                        select(Item).where(Item.project_id == project.id, Item.key == key, Item.deleted_at.is_(None))
                    ).scalar_one_or_none()
                if existing is None and not type_key:
                    raise ValidationFailed("column 'type' is required for new items")
                t = items_svc.item_type(project, existing.type_key if existing else type_key)
                field_values: dict[str, Any] = {}
                for f in t.fields:
                    if f.key in row and row[f.key] != "":
                        val: Any = row[f.key]
                        if f.type in ("multienum", "tags"):
                            val = [v.strip() for v in val.split(",") if v.strip()]
                        field_values[f.key] = val
                if existing:
                    change = items_svc.ContentChange(
                        title=row.get("title") or None,
                        description=row.get("description"),
                        fields=field_values or None,
                        applicability=row.get("applicability", items_svc._UNSET),
                    )
                    before = existing.version_no
                    items_svc.update_item(db, project, actor, existing, change, change_note="CSV/Excel import")
                    item = existing
                    action = "updated" if existing.version_no != before else "unchanged"
                else:
                    item = items_svc.create_item(
                        db, project, actor, type_key=t.key, title=row.get("title", ""),
                        description=row.get("description", ""), fields=field_values,
                        applicability=row.get("applicability") or None, kind="import",
                        change_note="CSV/Excel import",
                    )
                    action = "created"
                    if key:
                        created[key] = item
                if row.get("parent"):
                    pending_parents.append((idx, item, row["parent"].upper()))
                for lk in link_keys:
                    if row.get(lk):
                        pending_links.append((idx, item, lk, [k.strip().upper() for k in row[lk].split(",") if k]))
                results.append({"row": idx, "action": action, "key": item.key})
            except DomainError as exc:
                errors += 1
                results.append({"row": idx, "action": "error", "error": exc.message, "details": exc.details})

        def resolve(ref: str) -> Item | None:
            if ref in created:
                return created[ref]
            return db.execute(
                select(Item).where(Item.project_id == project.id, Item.key == ref, Item.deleted_at.is_(None))
            ).scalar_one_or_none()

        for idx, item, parent_key in pending_parents:
            parent = resolve(parent_key)
            try:
                if parent is None:
                    raise ValidationFailed(f"parent '{parent_key}' not found")
                structural = items_svc._structural_type_for(project, parent.type_key, item.type_key)
                already = db.execute(select(Link.id).where(Link.source_id == parent.id, Link.target_id == item.id,
                                                           Link.type_key == structural,
                                                           Link.deleted_at.is_(None))).scalar()
                if not already:
                    links_svc.create_link(db, project, actor, type_key=structural, source=parent, target=item)
            except DomainError as exc:
                errors += 1
                results.append({"row": idx, "action": "error", "error": exc.message})
        for idx, item, lk, targets in pending_links:
            for ref in targets:
                target = resolve(ref)
                try:
                    if target is None:
                        raise ValidationFailed(f"{lk} target '{ref}' not found")
                    already = db.execute(select(Link.id).where(Link.source_id == item.id, Link.target_id == target.id,
                                                               Link.type_key == lk,
                                                               Link.deleted_at.is_(None))).scalar()
                    if not already:
                        links_svc.create_link(db, project, actor, type_key=lk, source=item, target=target)
                except DomainError as exc:
                    errors += 1
                    results.append({"row": idx, "action": "error", "error": exc.message})
        if dry_run or errors:
            savepoint.rollback()
        else:
            savepoint.commit()
    except Exception:
        savepoint.rollback()
        raise
    summary: dict[str, int] = {}
    for r in results:
        summary[r["action"]] = summary.get(r["action"], 0) + 1
    return {"dry_run": dry_run, "applied": not dry_run and errors == 0, "rows": len(rows), "summary": summary,
            "results": results}
