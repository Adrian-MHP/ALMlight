"""Engineering-backbone features: computed fields, template updates, test execution, risks, BOM, planning,
notifications, saved queries and single sign-on."""

import base64
import csv
import hashlib
import io
import json
import time
import urllib.parse
import uuid

import jwt
import pytest
from conftest import new_item

from almlight.core.config import get_settings


def _link(api, project, type_, source, target, **extra):
    return api.post(f"/projects/{project}/links", {"type": type_, "source": source, "target": target, **extra},
                    expect=201)


# --- computed fields & template updates --------------------------------------------------------------------------


def test_risk_fields_are_computed_and_bounded(team, project):
    c = team["contributor"]
    r = new_item(c, project, "risk", "Pressure sensor drift", fields={"severity": 8, "occurrence": 5, "detection": 6,
                                                                      "rpn": 1, "risk_level": "low"})
    assert r["fields"]["rpn"] == 240 and r["fields"]["risk_level"] == "high"  # input for computed fields is ignored
    r2 = c.patch(f"/projects/{project}/items/{r['key']}", {"fields": {"occurrence": 2, "severity_after": 8,
                                                                      "occurrence_after": 1, "detection_after": 3}})
    assert r2["fields"]["rpn"] == 96 and r2["fields"]["risk_level"] == "medium"
    assert r2["fields"]["rpn_after"] == 24 and r2["fields"]["residual_level"] == "low"
    bad = c.patch(f"/projects/{project}/items/{r['key']}", {"fields": {"severity": 11}}, expect=422)
    assert "at most 10" in json.dumps(bad)


def test_template_updates_are_additive(admin, project):
    cfg = admin.get(f"/projects/{project}/config")["config"]
    assert admin.get(f"/projects/{project}/template-updates")["count"] == 0  # created from the current template
    old = json.loads(json.dumps(cfg))
    old["item_types"] = [t for t in old["item_types"] if t["key"] not in ("risk", "release", "hazard")]
    old["link_types"] = [lt for lt in old["link_types"]
                         if lt["key"] not in ("mitigated_by", "risk_for", "planned_in", "part_of", "depends_on")]
    old["workflows"] = [w for w in old["workflows"] if w["key"] not in ("risk", "release")]
    tc = next(t for t in old["item_types"] if t["key"] == "test_case")
    tc["fields"] = [f for f in tc["fields"] if f["key"] != "test_steps"]
    contains = next(lt for lt in old["link_types"] if lt["key"] == "contains")
    contains["attributes"] = [a for a in contains["attributes"] if a["key"] == "quantity"]
    admin.put(f"/projects/{project}/config", {"config": old})
    diff = admin.get(f"/projects/{project}/template-updates")
    assert {t["key"] for t in diff["item_types"]} == {"risk", "release", "hazard"}
    assert {lt["key"] for lt in diff["link_types"]} == {"mitigated_by", "risk_for", "planned_in", "part_of",
                                                        "depends_on"}
    assert {w["key"] for w in diff["workflows"]} == {"risk", "release"}
    assert {"type": "test_case", "key": "test_steps", "label": "Test steps", "field_type": "steps"} in diff["fields"]
    assert len(diff["link_attributes"]) == 4
    res = admin.post(f"/projects/{project}/template-updates")
    assert res["applied"]["count"] == diff["count"] and res["version"] > 1
    assert admin.get(f"/projects/{project}/template-updates")["count"] == 0
    new_item(admin, project, "risk", "Works after the update", fields={"severity": 2, "occurrence": 2, "detection": 2})


# --- test execution -------------------------------------------------------------------------------------------------


def test_step_by_step_test_execution_and_defect(team, project):
    c = team["contributor"]
    p = f"/projects/{project}"
    req = new_item(c, project, "system_requirement", "Measure pressure")
    tc = new_item(c, project, "test_case", "Full-scale test", fields={"test_steps": [
        {"step": "Apply 0 bar", "expected": "Output 4 mA"}, {"step": "Apply 10 bar", "expected": "Output 20 mA"},
        "Release pressure"]})
    assert tc["fields"]["test_steps"][2] == {"step": "Release pressure", "expected": ""}
    _link(c, project, "verifies", tc["key"], req["key"])
    run = new_item(c, project, "test_run", "Sample run")
    records = c.post(f"{p}/test-runs/{run['key']}/records", {"test_cases": [tc["key"]]})
    rec = records[0]
    assert [s["step"] for s in rec["steps"]] == ["Apply 0 bar", "Apply 10 bar", "Release pressure"]
    partial = c.patch(f"{p}/test-runs/{run['key']}/records/{rec['id']}",
                      {"step_results": [{"result": "passed", "actual": "4.0 mA"}]})[0]
    assert partial["result"] == "not_run" and len(partial["step_results"]) == 3  # in progress
    done = c.patch(f"{p}/test-runs/{run['key']}/records/{rec['id']}", {"step_results": [
        {"result": "passed", "actual": "4.0 mA"}, {"result": "failed", "actual": "19.2 mA"},
        {"result": "passed"}]})[0]
    assert done["result"] == "failed" and done["step_results"][1]["actual"] == "19.2 mA"
    bad = c.patch(f"{p}/test-runs/{run['key']}/records/{rec['id']}", {"step_results": [{"result": "great"}]},
                  expect=422)
    assert "result must be one of" in json.dumps(bad)
    res = c.post(f"{p}/test-runs/{run['key']}/records/{rec['id']}/defect", {"step": 2}, expect=201)
    defect = c.get(f"{p}/items/{res['defect']}")
    assert defect["type_key"] == "problem_report" and "19.2 mA" in defect["description"]
    assert res["records"][0]["defect"] == res["defect"]
    links = c.get(f"{p}/items/{res['defect']}/links")
    assert any(lk["other"]["key"] == tc["key"] for lk in links)
    again = c.post(f"{p}/test-runs/{run['key']}/records/{rec['id']}/defect", {"step": 2}, expect=409)
    assert again["error"]["details"]["defect"] == res["defect"]  # one defect per record, no duplicates
    team["viewer"].post(f"{p}/test-runs/{run['key']}/records/{rec['id']}/defect", {"step": 2}, expect=403)
    status = c.get(f"{p}/items/{req['key']}/verification")
    assert status["status"] == "failed"


# --- risks -----------------------------------------------------------------------------------------------------


def test_risk_matrix_register_and_mitigation(team, project):
    c = team["contributor"]
    p = f"/projects/{project}"
    high = new_item(c, project, "risk", "Overpressure bursts the cell", fields={"severity": 9, "occurrence": 5,
                                                                                "detection": 5})
    new_item(c, project, "risk", "Display flicker", fields={"severity": 2, "occurrence": 3, "detection": 2})
    overview = c.get(f"{p}/risks")
    assert overview["enabled"] and overview["summary"]["high"] == 1 and overview["summary"]["high_unmitigated"] == 1
    assert overview["matrix"][8][4] == 1 and overview["risks"][0]["key"] == high["key"]  # sorted by RPN
    req = new_item(c, project, "system_requirement", "Burst pressure 4x range")
    _link(c, project, "mitigated_by", high["key"], req["key"])
    after = c.get(f"{p}/risks")
    assert after["summary"]["high_unmitigated"] == 0 and after["risks"][0]["mitigations"] == 1


# --- BOM / where-used ----------------------------------------------------------------------------------------


def test_bom_quantities_effectivity_and_where_used(team, project, admin):
    c = team["contributor"]
    p = f"/projects/{project}"
    top = new_item(c, project, "system_element", "Transmitter", fields={"level": "product"})
    board = new_item(c, project, "system_element", "Electronics board", fields={"level": "module"})
    cap = new_item(c, project, "system_element", "Capacitor 10 uF", fields={"level": "component"})
    old = new_item(c, project, "system_element", "Old connector", fields={"level": "component"})
    _link(c, project, "contains", top["key"], board["key"], attributes={"quantity": 2, "find_number": "10"})
    _link(c, project, "contains", board["key"], cap["key"], attributes={"quantity": 3, "unit": "pcs"})
    _link(c, project, "contains", board["key"], old["key"], attributes={"quantity": 1, "effective_to": "2026-06-30"})
    bom = c.get(f"{p}/items/{top['key']}/bom")
    rows = {r["key"]: r for r in bom["rows"]}
    assert rows[cap["key"]]["total_quantity"] == 6 and rows[cap["key"]]["level"] == 2
    assert rows[board["key"]]["find_number"] == "10" and rows[board["key"]]["has_children"]
    later = c.get(f"{p}/items/{top['key']}/bom?effective=2026-12-01")
    assert old["key"] not in {r["key"] for r in later["rows"]}
    csv_text = c.get(f"{p}/items/{top['key']}/bom.csv").content.decode("utf-8-sig")
    assert cap["key"] in csv_text and "total_quantity" in csv_text.splitlines()[0]
    assert c.get(f"{p}/items/{top['key']}/bom.xlsx").content[:2] == b"PK"
    used = c.get(f"{p}/items/{cap['key']}/where-used")
    assert used["direct_parents"] == [board["key"]]
    assert [s["key"] for s in used["paths"][0]] == [top["key"], board["key"]]


# --- planning ------------------------------------------------------------------------------------------------


def test_release_plan_progress_and_burnup(team, project):
    c, m = team["contributor"], team["manager"]
    p = f"/projects/{project}"
    rel = new_item(c, project, "release", "Release 1.0", fields={"start_date": "2026-09-01",
                                                                 "due_date": "2027-03-31"})
    reqs = [new_item(c, project, "system_requirement", f"R{i}") for i in range(4)]
    for r in reqs:
        _link(c, project, "planned_in", r["key"], rel["key"])
    risk = new_item(c, project, "test_run", "Qualification run")
    _link(c, project, "planned_in", risk["key"], rel["key"])
    for r in reqs[:2]:
        c.post(f"{p}/items/{r['key']}/transitions/submit", {})
        m.post(f"{p}/items/{r['key']}/transitions/approve", {"password": "password-123"}, expect=(200, 201))
    plans = c.get(f"{p}/plans")
    assert plans["enabled"] and plans["releases"][0]["key"] == rel["key"]
    detail = c.get(f"{p}/plans/{rel['key']}")
    assert detail["total"] == 5
    assert len(detail["items"]) == 5 and detail["burnup"][-1]["scope"] == 5
    listed = c.get(f"{p}/items?planned_in={rel['key']}&type=system_requirement")
    assert listed["total"] == 4


# --- notifications ----------------------------------------------------------------------------------------------


def _notes(api):
    return api.get("/notifications")


def test_notifications_watch_owner_approval_mention(team, project, app):
    c, r, m, v = team["contributor"], team["reviewer"], team["manager"], team["viewer"]
    p = f"/projects/{project}"
    spec = new_item(c, project, "specification", "Watched spec")
    req = new_item(c, project, "system_requirement", "Inside the spec", parent=spec["key"])
    v.put(f"{p}/items/{spec['key']}/watch")
    assert v.get(f"{p}/items/{spec['key']}/watch") == {"watching": True, "watchers": 1}
    m.patch(f"{p}/items/{req['key']}", {"description": "Changed by the manager"})
    vn = _notes(v)
    assert vn["unread"] == 1 and vn["items"][0]["kind"] == "watch" and vn["items"][0]["item_key"] == req["key"]
    cn = _notes(c)
    assert cn["items"][0]["kind"] == "owner" and "Edited" in cn["items"][0]["title"]
    assert _notes(m)["unread"] == 0  # the actor is never notified
    c.post(f"{p}/items/{req['key']}/transitions/submit", {})
    c.post(f"{p}/items/{req['key']}/transitions/approve", {})
    assert any(n["kind"] == "approval" for n in _notes(r)["items"])
    c.post(f"{p}/items/{req['key']}/comments", {"body": "@rene can you check the range? cc @nobody",
                                                 "document": spec["key"]}, expect=201)
    mention = _notes(r)["items"][0]
    assert mention["kind"] == "mention" and "Commented" in mention["title"]
    assert mention["item_key"] == req["key"]  # opens the commented item, not the document it was made in
    r.post("/notifications/read", {"all": True})
    assert _notes(r)["unread"] == 0
    v.delete(f"{p}/items/{spec['key']}/watch")
    assert v.get(f"{p}/items/{spec['key']}/watch")["watching"] is False


def test_bulk_changes_collapse_and_dry_runs_notify_nobody(team, project):
    c, v = team["contributor"], team["viewer"]
    p = f"/projects/{project}"
    spec = new_item(c, project, "specification", "Import target")
    v.put(f"{p}/items/{spec['key']}/watch")
    rows = "key,type,title,parent\n" + "\n".join(f"N{i},system_requirement,Imported {i},{spec['key']}"
                                                 for i in range(12))
    c.post(f"{p}/import", files={"file": ("x.csv", io.BytesIO(rows.encode()), "text/csv")}, data={"dry_run": "true"})
    assert _notes(v)["unread"] == 0  # dry run: rolled back, nobody notified
    c.post(f"{p}/import", files={"file": ("x.csv", io.BytesIO(rows.encode()), "text/csv")}, data={"dry_run": "false"})
    notes = _notes(v)
    assert notes["unread"] == 1 and notes["items"][0]["kind"] == "digest"
    assert "changes in" in notes["items"][0]["title"]
    # a re-import that only adds links reports those rows as updated, not unchanged
    rel = new_item(c, project, "release", "Release 9")
    keys = [i["key"] for i in c.get(f"{p}/items?type=system_requirement&q=Imported&limit=50")["items"]]
    plan = "key,planned_in\n" + "\n".join(f"{k},{rel['key']}" for k in keys)
    res = c.post(f"{p}/import", files={"file": ("p.csv", io.BytesIO(plan.encode()), "text/csv")}, data={"dry_run": "false"})
    assert res["summary"] == {"updated": 12, "links": 12}
    res = c.post(f"{p}/import", files={"file": ("p.csv", io.BytesIO(plan.encode()), "text/csv")}, data={"dry_run": "false"})
    assert res["summary"] == {"unchanged": 12}


def test_email_digest(team, project, db, monkeypatch):
    from almlight.modules.identity.models import User
    from almlight.modules.notifications import service as notif

    c, v = team["contributor"], team["viewer"]
    s = get_settings()
    monkeypatch.setattr(s, "smtp_host", "smtp.example.com")
    monkeypatch.setattr(s, "smtp_from", "alm@example.com")
    monkeypatch.setattr(s, "public_url", "https://alm.example.com")
    sent = []
    monkeypatch.setattr(notif, "_send", lambda msg: sent.append(msg))
    user = db.get(User, uuid.UUID(team["users"]["viewer"]["id"]))
    user.email = "vic@example.com"
    db.commit()
    item = new_item(c, project, "system_requirement", "Mail me")
    v.put(f"/projects/{project}/items/{item['key']}/watch")
    c.patch(f"/projects/{project}/items/{item['key']}", {"title": "Mail me (changed)"})
    assert notif.send_digests(db, min_age_seconds=0) == 1
    body = sent[0].get_content()
    assert sent[0]["To"] == "vic@example.com" and f"https://alm.example.com/p/{project}/items/{item['key']}" in body
    assert notif.send_digests(db, min_age_seconds=0) == 0  # nothing twice
    v.put("/notifications/preferences", {"email_notifications": False})


# --- saved queries -----------------------------------------------------------------------------------------------


def test_saved_queries_and_traceability_gaps(team, project):
    c, v, m = team["contributor"], team["viewer"], team["manager"]
    p = f"/projects/{project}"
    covered = new_item(c, project, "system_requirement", "Covered")
    new_item(c, project, "system_requirement", "Not covered")
    tc = new_item(c, project, "test_case", "A test")
    _link(c, project, "verifies", tc["key"], covered["key"])
    gaps = c.get(f"{p}/items?type=system_requirement&link_type=verifies&link_dir=in&link_exists=false")
    assert [i["title"] for i in gaps["items"]] == ["Not covered"]
    q = {"type": ["system_requirement"], "link_type": "verifies", "link_dir": "in", "link_exists": False}
    shared = c.post(f"{p}/queries", {"name": "Unverified requirements", "query": q, "shared": True}, expect=201)
    mine = v.post(f"{p}/queries", {"name": "My drafts", "query": {"state": ["draft"]}}, expect=201)
    v.post(f"{p}/queries", {"name": "x", "query": {}, "shared": True}, expect=403)  # viewers cannot share
    c.post(f"{p}/queries", {"name": "bad", "query": {"evil": 1}}, expect=422)
    listed = {x["name"]: x for x in v.get(f"{p}/queries")}
    assert listed["Unverified requirements"]["count"] == 1 and listed["My drafts"]["count"] >= 2
    assert "My drafts" not in {x["name"] for x in c.get(f"{p}/queries")}  # personal stays personal
    assert listed["Unverified requirements"]["can_edit"] is False and listed["My drafts"]["can_edit"] is True
    assert {x["name"]: x["can_edit"] for x in m.get(f"{p}/queries")}["Unverified requirements"] is True
    v.put(f"{p}/queries/{shared['id']}", {"name": "hijack", "query": q, "shared": True}, expect=403)
    v.delete(f"{p}/queries/{shared['id']}", expect=403)
    m.put(f"{p}/queries/{shared['id']}", {"name": "Unverified (CCB)", "query": q, "shared": True})
    v.delete(f"{p}/queries/{mine['id']}")


def test_export_of_the_current_view(team, project):
    """Export with the item list's filter delivers exactly that view, across all pages; without, the whole type."""
    c = team["contributor"]
    p = f"/projects/{project}"
    covered = new_item(c, project, "system_requirement", "Export covered")
    new_item(c, project, "system_requirement", "Export gap one")
    new_item(c, project, "system_requirement", "Export gap two")
    tc = new_item(c, project, "test_case", "Export test")
    _link(c, project, "verifies", tc["key"], covered["key"])
    gap = "type=system_requirement&link_type=verifies&link_dir=in&link_exists=false"
    rows = list(csv.reader(io.StringIO(c.get(f"{p}/export.csv?{gap}").content.decode("utf-8-sig"))))
    titles = {r[2] for r in rows[1:]}
    assert {"Export gap one", "Export gap two"} <= titles and "Export covered" not in titles and "Export test" not in titles
    everything = c.get(f"{p}/export.csv?type=system_requirement").content.decode("utf-8-sig")
    assert "Export covered" in everything and "Export test" not in everything
    searched = c.get(f"{p}/export.xlsx?q=gap%20two")
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(searched.content)).active
    assert [r[2] for r in ws.iter_rows(min_row=2, values_only=True)] == ["Export gap two"]
    c.get(f"{p}/export.csv?link_dir=sideways", expect=422)


# --- single sign-on ------------------------------------------------------------------------------------------------


@pytest.fixture
def fake_idp(monkeypatch):
    from cryptography.hazmat.primitives.asymmetric import rsa

    from almlight.modules.identity import oidc

    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    jwk = json.loads(jwt.algorithms.RSAAlgorithm.to_jwk(key.public_key()))
    jwk.update({"kid": "k1", "use": "sig", "alg": "RS256"})
    issuer = "https://idp.example.com/realms/eng"
    s = get_settings()
    for k, val in (("oidc_issuer", issuer), ("oidc_client_id", "ariadne"), ("oidc_client_secret", "s3cret"),
                   ("oidc_admin_group", "alm-admins")):
        monkeypatch.setattr(s, k, val)
    oidc._cache.clear()
    state = {"claims": {}, "nonce": None, "verifier": None}

    def http_get(url):
        if url.endswith("/.well-known/openid-configuration"):
            return {"issuer": issuer, "authorization_endpoint": f"{issuer}/auth", "token_endpoint": f"{issuer}/token",
                    "jwks_uri": f"{issuer}/certs"}
        if url.endswith("/certs"):
            return {"keys": [jwk]}
        raise AssertionError(url)

    def http_post(url, data, auth):
        assert auth == ("ariadne", "s3cret") and data["grant_type"] == "authorization_code"
        challenge = base64.urlsafe_b64encode(hashlib.sha256(data["code_verifier"].encode()).digest()).rstrip(b"=")
        assert challenge.decode() == state["challenge"]  # PKCE
        now = int(time.time())
        claims = {"iss": issuer, "aud": "ariadne", "sub": "u-1", "iat": now, "exp": now + 300,
                  "nonce": state["nonce"], **state["claims"]}
        return {"id_token": jwt.encode(claims, key, algorithm="RS256", headers={"kid": "k1"})}

    monkeypatch.setattr(oidc, "_http_get", http_get)
    monkeypatch.setattr(oidc, "_http_post", http_post)
    yield state
    oidc._cache.clear()


def test_single_sign_on(app, fake_idp, anon):
    from fastapi.testclient import TestClient

    assert anon.get("/auth/oidc/config")["enabled"] is True
    client = TestClient(app)
    r = client.get("/api/v1/auth/oidc/login?next=/p/X", follow_redirects=False)
    assert r.status_code == 302
    q = urllib.parse.parse_qs(urllib.parse.urlparse(r.headers["location"]).query)
    assert q["code_challenge_method"] == ["S256"] and q["client_id"] == ["ariadne"]
    fake_idp.update(nonce=q["nonce"][0], challenge=q["code_challenge"][0],
                    claims={"preferred_username": "erika", "name": "Erika Engineer", "email": "erika@example.com",
                            "groups": ["alm-admins"]})
    bad = client.get("/api/v1/auth/oidc/callback?code=abc&state=wrong", follow_redirects=False)
    assert bad.status_code == 401
    done = client.get(f"/api/v1/auth/oidc/callback?code=abc&state={q['state'][0]}", follow_redirects=False)
    assert done.status_code == 302 and done.headers["location"] == "/p/X"
    me = client.get("/api/v1/auth/me").json()
    assert me["username"] == "erika" and me["display_name"] == "Erika Engineer" and me["is_admin"] is True
    # second sign-in without the admin group: same user, admin removed
    r2 = client.get("/api/v1/auth/oidc/login", follow_redirects=False)
    q2 = urllib.parse.parse_qs(urllib.parse.urlparse(r2.headers["location"]).query)
    fake_idp.update(nonce="tampered", challenge=q2["code_challenge"][0], claims={"preferred_username": "erika"})
    assert client.get(f"/api/v1/auth/oidc/callback?code=x&state={q2['state'][0]}",
                      follow_redirects=False).status_code == 401  # nonce mismatch
    r3 = client.get("/api/v1/auth/oidc/login", follow_redirects=False)
    q3 = urllib.parse.parse_qs(urllib.parse.urlparse(r3.headers["location"]).query)
    fake_idp.update(nonce=q3["nonce"][0], challenge=q3["code_challenge"][0], claims={"preferred_username": "erika"})
    client.get(f"/api/v1/auth/oidc/callback?code=x&state={q3['state'][0]}", follow_redirects=False)
    assert client.get("/api/v1/auth/me").json()["is_admin"] is False
