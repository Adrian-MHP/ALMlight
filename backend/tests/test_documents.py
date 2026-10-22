"""Live-document features: time travel, history, compare, Word/ReqIF round trip, comments, recycle bin, collections."""

import io
import time
import zipfile

from conftest import approve, new_item
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from test_reqif_export_reports import DOORS_REQIF, _import


def _spec(c, project):
    spec = new_item(c, project, "specification", "System spec")
    h1 = new_item(c, project, "heading", "Measurement", parent=spec["key"])
    r1 = new_item(c, project, "system_requirement", "Measure pressure", parent=h1["key"],
                  description="Range 0-10 bar.", fields={"priority": "must"})
    r2 = new_item(c, project, "system_requirement", "Accuracy", parent=h1["key"], description="0.1 % of span")
    h2 = new_item(c, project, "heading", "Interfaces", parent=spec["key"])
    r3 = new_item(c, project, "system_requirement", "Current output", parent=h2["key"], applicability="OUT_420")
    return spec, h1, r1, r2, h2, r3


def _doc(c, project, key, **params):
    from urllib.parse import urlencode

    qs = urlencode(params)
    return c.get(f"/projects/{project}/items/{key}/document{'?' + qs if qs else ''}")


def _titles(node):
    return [(n["display_title"], [x["title"] for x in n["children"]]) for n in node["children"]]


def _now(c):
    # the server clock decides; take the newest audit timestamp as "now"
    return c.get("/projects/PX/audit?limit=1")[0]["ts"]


def test_document_time_travel(team, project):
    c = team["contributor"]
    p = f"/projects/{project}"
    spec, h1, r1, r2, h2, r3 = _spec(c, project)
    live = _doc(c, project, spec["key"])
    assert _titles(live["root"]) == [("1 Measurement", ["Measure pressure", "Accuracy"]),
                                     ("2 Interfaces", ["Current output"])]
    t1 = _now(c)
    time.sleep(0.05)
    # Change content, order, structure and lifecycle after t1.
    c.patch(f"{p}/items/{r1['key']}", {"title": "Measure process pressure", "description": "Range 0-16 bar."})
    c.post(f"{p}/items/{r2['key']}/move", {"parent": h1["key"], "position": 0.5})
    new_item(c, project, "system_requirement", "Diagnostics", parent=h2["key"])
    c.delete(f"{p}/links/{_doc(c, project, spec['key'])['root']['children'][1]['children'][0]['link_id']}")
    c.post(f"{p}/items/{r1['key']}/transitions/submit", {})

    now = _doc(c, project, spec["key"])
    assert _titles(now["root"]) == [("1 Measurement", ["Accuracy", "Measure process pressure"]),
                                    ("2 Interfaces", ["Diagnostics"])]
    past = _doc(c, project, spec["key"], at=t1)
    assert past["mode"] == "at"
    assert _titles(past["root"]) == [("1 Measurement", ["Measure pressure", "Accuracy"]),
                                     ("2 Interfaces", ["Current output"])]
    old_r1 = past["root"]["children"][0]["children"][0]
    assert old_r1["description"] == "<p>Range 0-10 bar.</p>" and old_r1["state_key"] == "draft" and old_r1["version"] == "A.1"
    assert now["root"]["children"][0]["children"][1]["state_key"] == "in_review"
    c.get(f"{p}/items/{spec['key']}/document?at=2000-01-01T00:00:00Z", expect=404)
    c.get(f"{p}/items/{spec['key']}/document?at=yesterday", expect=422)
    # Export the past state as Word.
    old = Document(io.BytesIO(c.get(f"{p}/items/{spec['key']}/export.docx?at={t1}").content))
    text = "\n".join(par.text for par in old.paragraphs)
    assert "Current output" in text and "Diagnostics" not in text and "1 Measurement" in text


def test_document_baseline_variant_and_history(team, project):
    c, r = team["contributor"], team["reviewer"]
    p = f"/projects/{project}"
    spec, h1, r1, r2, h2, r3 = _spec(c, project)
    bl = r.post(f"{p}/baselines", {"name": "Draft 1", "scope": spec["key"]})
    c.patch(f"{p}/items/{r2['key']}", {"title": "Accuracy 0.1 %"})
    at_bl = _doc(c, project, spec["key"], baseline=bl["id"])
    assert at_bl["mode"] == "baseline" and at_bl["baseline"]["name"] == "Draft 1"
    assert at_bl["root"]["children"][0]["children"][1]["title"] == "Accuracy"
    new_item(c, project, "feature", "IO-Link", fields={"code": "OUT_IOLINK"})
    var = new_item(c, project, "variant", "IO-Link variant", fields={"selection": ["OUT_IOLINK"]})
    filtered = _doc(c, project, spec["key"], variant=var["key"])
    assert filtered["variant"]["key"] == var["key"]
    assert filtered["root"]["children"][1]["children"] == []  # 4-20 mA requirement not applicable
    hist = c.get(f"{p}/items/{spec['key']}/document-history")
    actions = {e["action"] for e in hist["entries"]}
    assert {"item.created", "link.created", "item.updated", "baseline.created"} <= actions
    assert hist["baselines"][0]["name"] == "Draft 1" and hist["baselines"][0]["scope_is_document"] is True
    assert any(r2["key"] in e["items"] for e in hist["entries"] if e["action"] == "item.updated")


def test_compare_revisions_branches_and_documents(team, project, admin):
    c, r = team["contributor"], team["reviewer"]
    p = f"/projects/{project}"
    spec, h1, r1, r2, h2, r3 = _spec(c, project)
    bl = r.post(f"{p}/baselines", {"name": "B1", "scope": spec["key"]})
    c.patch(f"{p}/items/{r1['key']}", {"description": "Range 0-16 bar.", "fields": {"priority": "should"}})
    c.post(f"{p}/items/{r2['key']}/move", {"parent": h1["key"], "position": 0.5})
    added = new_item(c, project, "system_requirement", "Diagnostics", parent=h2["key"])
    c.delete(f"{p}/links/{_doc(c, project, spec['key'])['root']['children'][1]['children'][0]['link_id']}")
    res = c.post(f"{p}/compare-documents", {"left": {"document": spec["key"], "baseline": bl["id"]},
                                            "right": {"document": spec["key"]}})
    by_key = {(row["right"] or row["left"])["key"]: row for row in res["rows"]}
    assert by_key[r1["key"]]["status"] == "changed"
    desc = next(ch for ch in by_key[r1["key"]]["changes"] if ch["field"] == "description")
    assert {"op": "delete", "text": "10"} in desc["diff"] and {"op": "insert", "text": "16"} in desc["diff"]
    assert any(ch["field"] == "fields.priority" for ch in by_key[r1["key"]]["changes"])
    assert by_key[added["key"]]["status"] == "added" and by_key[r3["key"]]["status"] == "removed"
    assert by_key[r1["key"]]["moved"] or by_key[r2["key"]]["moved"]
    assert res["summary"]["added"] == 1 and res["summary"]["removed"] == 1
    only_title = c.post(f"{p}/compare-documents", {"left": {"document": spec["key"], "baseline": bl["id"]},
                                                   "right": {"document": spec["key"]}, "fields": ["kind"]})
    row = next(x for x in only_title["rows"] if (x["right"] or {}).get("key") == r1["key"])
    assert not any(ch["field"] == "fields.priority" for ch in row["changes"])

    # Branch vs master: copies are paired with their sources.
    br = c.post(f"{p}/items/{spec['key']}/branch", {"name": "Variant X"}, expect=201)
    copy_key = c.get(f"{p}/items/{br['branch_root']}/document")["root"]["children"][0]["children"][0]["key"]
    c.patch(f"{p}/items/{copy_key}", {"title": "Accuracy for variant X"})
    vs = c.post(f"{p}/compare-documents", {"left": {"document": spec["key"]}, "right": {"document": br["branch_root"]}})
    assert vs["summary"]["changed"] == 1 and vs["summary"]["added"] == 0 and vs["summary"]["removed"] == 0

    # Two documents in different projects (paired by title).
    admin.post("/projects", {"key": "OTHER", "name": "Other", "template": "systems_engineering"})
    admin.put("/projects/OTHER/members", {"user_id": team["users"]["contributor"]["id"], "role": "viewer"})
    o_spec = new_item(admin, "OTHER", "specification", "Other spec")
    o_h = new_item(admin, "OTHER", "heading", "Measurement", parent=o_spec["key"])
    new_item(admin, "OTHER", "system_requirement", "Accuracy", parent=o_h["key"], description="0.2 % of span")
    cross = c.post(f"{p}/compare-documents", {"left": {"document": spec["key"]},
                                              "right": {"project": "OTHER", "document": o_spec["key"]}})
    acc = next(x for x in cross["rows"] if x["right"] and x["right"]["title"] == "Accuracy")
    assert acc["status"] == "changed" and acc["left"]["key"] == r2["key"]


def _set_sdt_text(doc, tag_part, item_id, new_text, desc=False):
    for sdt in doc.element.body.iter(qn("w:sdt")):
        tag = sdt.find(f"{qn('w:sdtPr')}/{qn('w:tag')}").get(qn("w:val"))
        if tag.startswith(f"alm:{item_id}:") and tag.endswith(f":{tag_part}"):
            ts = list(sdt.iter(qn("w:t")))
            ts[0].text = new_text
            for t in ts[1:]:
                t.text = ""
            return sdt
    raise AssertionError(f"no control for {item_id} {tag_part}")


def test_word_round_trip(team, project):
    c, r = team["contributor"], team["reviewer"]
    p = f"/projects/{project}"
    spec, h1, r1, r2, h2, r3 = _spec(c, project)
    c.patch(f"{p}/items/{r2['key']}", {"description": "Accuracy of **0.1 %** of span.\n\n- at 20 degC\n- at 50 degC"})
    exported = c.get(f"{p}/items/{spec['key']}/export.docx?roundtrip=edit")
    assert exported.headers["content-disposition"].endswith('review-edit.docx"')
    doc = Document(io.BytesIO(exported.content))
    assert doc.core_properties.identifier == f"almlight:{project}:{spec['key']}"
    tags = [s.find(f"{qn('w:sdtPr')}/{qn('w:tag')}").get(qn("w:val")) for s in doc.element.body.iter(qn("w:sdt"))]
    assert len(tags) == 2 * 3 + 2  # title+text for 3 requirements, title for 2 headings

    def upload(content, dry_run=True, key=spec["key"], expect=200):
        return c.post(f"{p}/items/{key}/import-docx", files={"file": ("review.docx", io.BytesIO(content), "application/x")},
                      data={"dry_run": str(dry_run).lower()}, expect=expect)

    untouched = upload(exported.content)
    assert untouched["summary"]["updated"] == 0 and untouched["summary"]["unchanged"] == 5  # formatting survives

    # Reviewer edits r1's text and r3's title and comments on r1.
    sdt = _set_sdt_text(doc, "desc", r1["id"], "Range 0-25 bar.")
    _set_sdt_text(doc, "title", r3["id"], "Current output 4-20 mA")
    para = Paragraph(next(sdt.iter(qn("w:p"))), doc._body)
    doc.add_comment(para.runs, text="Is 25 bar confirmed by the customer?", author="Jane Reviewer", initials="JR")
    buf = io.BytesIO()
    doc.save(buf)
    edited = buf.getvalue()

    dry = upload(edited)
    assert dry["applied"] is False and dry["summary"]["updated"] == 2 and dry["summary"]["comments"] == 1
    assert c.get(f"{p}/items/{r1['key']}")["description"] == "<p>Range 0-10 bar.</p>"
    res = upload(edited, dry_run=False)
    assert res["applied"] is True
    after = c.get(f"{p}/items/{r1['key']}")
    assert after["description"] == "<p>Range 0-25 bar.</p>" and after["version"] == "A.2"
    assert c.get(f"{p}/items/{r3['key']}")["title"] == "Current output 4-20 mA"
    assert c.get(f"{p}/items/{r2['key']}")["version"] == "A.2"  # untouched: still the exported version
    comments = c.get(f"{p}/items/{r1['key']}/comments")
    assert comments[0]["external_author"] == "Jane Reviewer" and comments[0]["source"] == "word"
    again = upload(edited, dry_run=False)  # same file again: no duplicates, edits already applied
    assert again["summary"]["comments"] == 0 and again["summary"]["conflicts"] == 2
    assert len(c.get(f"{p}/items/{r1['key']}/comments")) == 1

    # Conflict: the item changed in ALMlight after the export; frozen items are reported.
    exp2 = Document(io.BytesIO(c.get(f"{p}/items/{spec['key']}/export.docx?roundtrip=edit").content))
    c.patch(f"{p}/items/{r1['key']}", {"description": "Changed in ALMlight meanwhile."})
    _set_sdt_text(exp2, "desc", r1["id"], "Changed by the reviewer.")
    c.post(f"{p}/items/{r2['key']}/transitions/submit", {})
    _set_sdt_text(exp2, "desc", r2["id"], "Reviewer text for a frozen item.")
    buf = io.BytesIO()
    exp2.save(buf)
    res = upload(buf.getvalue(), dry_run=False)
    statuses = {i["key"]: i["status"] for i in res["items"]}
    assert statuses[r1["key"]] == "conflict" and statuses[r2["key"]] == "frozen"
    assert c.get(f"{p}/items/{r1['key']}")["description"] == "<p>Changed in ALMlight meanwhile.</p>"

    # Wrong document / not a round-trip export / comments-only mode.
    other = new_item(c, project, "specification", "Other")
    upload(edited, key=other["key"], expect=422)
    plain = c.get(f"{p}/items/{spec['key']}/export.docx").content
    upload(plain, expect=422)
    upload(b"not a zip", expect=422)
    review = Document(io.BytesIO(c.get(f"{p}/items/{spec['key']}/export.docx?roundtrip=comments").content))
    prot = review.settings.element.find(qn("w:documentProtection"))
    assert prot is not None and prot.get(qn("w:edit")) == "comments"
    c.get(f"{p}/items/{spec['key']}/export.docx?roundtrip=edit&at=2026-01-01T00:00:00Z", expect=422)
    approve(r2["key"], project, c, r)


def test_reqif_export_keeps_partner_identifiers(team, project):
    c = team["contributor"]
    p = f"/projects/{project}"
    _import(c, project, DOORS_REQIF.encode(), dry_run=False,
            mapping={"types": {"Customer Requirement": "stakeholder_requirement"}, "relations": {"refines": "relates_to"}})
    spec = c.get(f"{p}/items?type=specification")["items"][0]
    xml = c.get(f"{p}/items/{spec['key']}/reqif").content.decode()
    assert 'IDENTIFIER="o-r1"' in xml and 'IDENTIFIER="o-r2"' in xml and 'IDENTIFIER="s1"' in xml
    again = _import(c, project, xml.encode(), dry_run=False)
    assert again["summary"]["created"] == 0


def test_comment_threads(team, project):
    c, r, m = team["contributor"], team["reviewer"], team["manager"]
    p = f"/projects/{project}"
    spec, h1, r1, *_ = _spec(c, project)
    url = f"{p}/items/{r1['key']}/comments"
    root = r.post(url, {"body": "Range unclear", "document": spec["key"]}, expect=201)
    reply = c.post(url, {"body": "Clarified with PM", "parent_id": root["id"]}, expect=201)
    c.post(url, {"body": "x", "parent_id": "00000000-0000-0000-0000-000000000000"}, expect=404)
    threads = c.get(url)
    assert len(threads) == 1 and threads[0]["replies"][0]["body"] == "Clarified with PM"
    assert threads[0]["document_key"] == spec["key"]
    c.post(f"{url}/{reply['id']}/resolve", expect=422)  # resolve the thread, not a reply
    c.post(f"{url}/{root['id']}/resolve")
    assert c.get(url)[0]["resolved"] is True and c.get(f"{url}?resolved=false") == []
    doc_comments = c.get(f"{p}/items/{spec['key']}/document-comments")
    assert doc_comments[0]["item_key"] == r1["key"]
    assert c.get(f"{p}/items/{spec['key']}/document")["root"]["children"][0]["children"][0]["comment_count"] == 0
    c.post(url, {"body": "Reopening", "parent_id": root["id"]}, expect=201)  # a reply reopens the thread
    assert c.get(url)[0]["resolved"] is False
    assert c.get(f"{p}/items/{spec['key']}/document")["root"]["children"][0]["children"][0]["comment_count"] == 1
    c.delete(f"{url}/{root['id']}", expect=403)
    m.delete(f"{url}/{root['id']}")
    assert c.get(url) == []


def test_recycle_bin(team, project):
    c, m = team["contributor"], team["manager"]
    p = f"/projects/{project}"
    spec, h1, r1, r2, *_ = _spec(c, project)
    stk = new_item(c, project, "stakeholder_requirement", "Need")
    c.post(f"{p}/links", {"type": "derives_from", "source": r2["key"], "target": stk["key"]})
    m.delete(f"{p}/items/{r2['key']}")
    assert [b["key"] for b in c.get(f"{p}/recycle-bin")] == [r2["key"]]
    c.post(f"{p}/items/{r2['key']}/undelete", expect=403)
    restored = m.post(f"{p}/items/{r2['key']}/undelete")
    assert restored["key"] == r2["key"]
    assert c.get(f"{p}/recycle-bin") == []
    titles = [n["title"] for n in c.get(f"{p}/items/{spec['key']}/document")["root"]["children"][0]["children"]]
    assert titles == ["Measure pressure", "Accuracy"]
    assert c.get(f"{p}/items/{r2['key']}/links")[0]["type"] == "derives_from"
    m.post(f"{p}/items/{r2['key']}/undelete", expect=409)


def test_collections(team, project):
    c, r, m = team["contributor"], team["reviewer"], team["manager"]
    p = f"/projects/{project}"
    spec, *_ = _spec(c, project)
    spec2 = new_item(c, project, "specification", "Test spec")
    new_item(c, project, "test_case", "Range test", parent=spec2["key"])
    c.post(f"{p}/collections", {"name": "Release 1.0"}, expect=403)
    col = r.post(f"{p}/collections", {"name": "Release 1.0", "description": "Customer delivery"}, expect=201)
    r.post(f"{p}/collections/{col['id']}/status", {"status": "closed"}, expect=422)  # empty
    col = r.post(f"{p}/collections/{col['id']}/documents", {"document": spec["key"]})
    assert col["entries"][0]["baseline"] == f"Release 1.0 — {spec['key']}"
    bl = r.post(f"{p}/baselines", {"name": "Test spec v1", "scope": spec2["key"]})
    wrong = r.post(f"{p}/baselines", {"name": "Only spec", "scope": spec["key"]})
    r.post(f"{p}/collections/{col['id']}/documents", {"document": spec2["key"], "baseline": wrong["id"]}, expect=422)
    r.post(f"{p}/collections/{col['id']}/documents", {"document": spec2["key"], "baseline": bl["id"]})
    r.post(f"{p}/collections/{col['id']}/documents", {"document": "SYS-1"}, expect=422)  # not a document
    col = r.post(f"{p}/collections/{col['id']}/status", {"status": "closed"})
    assert col["status"] == "closed" and len(col["entries"]) == 2
    r.post(f"{p}/collections/{col['id']}/documents", {"document": spec["key"]}, expect=409)
    r.post(f"{p}/collections/{col['id']}/status", {"status": "open"}, expect=409)  # only managers reopen
    m.delete(f"{p}/collections/{col['id']}", expect=409)
    z = zipfile.ZipFile(io.BytesIO(c.get(f"{p}/collections/{col['id']}/export.zip?format=pdf").content))
    names = z.namelist()
    assert "MANIFEST.txt" in names and sum(n.endswith(".pdf") for n in names) == 2
    m.post(f"{p}/collections/{col['id']}/status", {"status": "open"})
    m.delete(f"{p}/collections/{col['id']}")
    assert c.get(f"{p}/collections") == []
