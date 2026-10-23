"""Customer Word templates: upload checks, placeholders (also split over runs), content position, table of contents,
styles and list numbering copied into minimal templates, permissions."""

import io
import zipfile

from conftest import new_item
from docx import Document
from docx.oxml.ns import qn

RICH = ("<p>Applies to all variants.</p><ul><li><p>first point</p></li><li><p>second point</p></li></ul>"
        "<ol><li><p>step one</p></li><li><p>step two</p></li></ol>"
        "<table><tr><th><p>Signal</p></th><th><p>Range</p></th></tr><tr><td><p>Current</p></td>"
        "<td><p>4-20 mA</p></td></tr></table>")
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _minimal_template() -> bytes:
    """A corporate template without list, heading 2+ and grid table styles; a placeholder split over two runs;
    a closing page after the content marker."""
    doc = Document()
    styles = doc.styles.element
    for st in list(styles.findall(qn("w:style"))):
        name = st.find(qn("w:name"))
        if name is not None and (name.get(qn("w:val")).startswith(("List", "heading 2", "heading 3", "Light Grid"))):
            styles.remove(st)
    doc.sections[0].header.paragraphs[0].text = "ACME Corp. — {{project_name}}"
    p = doc.add_paragraph()
    p.add_run("{{ti").bold = True
    p.add_run("tle}} ({{key}} {{version}}, {{state}}) by {{author}} on {{date}} {{unknown_thing}}")
    doc.add_paragraph("{{toc}}")
    doc.add_paragraph("{{content}}")
    doc.add_paragraph("END OF DOCUMENT — signatures on the next page")
    return _bytes(doc)


def _upload(api, project, name, content, filename="corporate.docx", expect=201):
    return api.post(f"/projects/{project}/doc-templates", files={"file": (filename, io.BytesIO(content), DOCX)},
                    data={"name": name}, expect=expect)


def _texts(doc):
    return [p.text for p in doc.paragraphs]


def test_template_upload_rules_and_sample(team, project, admin):
    c, m = team["contributor"], team["manager"]
    p = f"/projects/{project}"
    sample = c.get(f"{p}/doc-templates/sample.docx")
    assert sample.headers["content-type"].startswith(DOCX)
    t = _upload(m, project, "Corporate", sample.content)
    assert t["has_content"] and t["has_toc"] and {"title", "version", "content"} <= set(t["placeholders"])
    _upload(c, project, "Mine", sample.content, expect=403)  # managers manage templates
    bad = _upload(m, project, "Text", b"not a word file", expect=422)
    assert "Not a Word file" in bad["error"]["message"]
    macro = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(sample.content)) as src, zipfile.ZipFile(macro, "w") as dst:
        for info in src.infolist():
            dst.writestr(info, src.read(info.filename))
        dst.writestr("word/vbaProject.bin", b"\x00")
    assert "macros" in _upload(m, project, "Macro", macro.getvalue(), expect=422)["error"]["message"]
    # a .dotx (Word template file) is accepted and stored as a document
    dotx = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(sample.content)) as src, zipfile.ZipFile(dotx, "w") as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(b"document.main+xml", b"template.main+xml")
            dst.writestr(info, data)
    as_dotx = _upload(m, project, "From dotx", dotx.getvalue(), filename="corporate.dotx")
    stored = c.get(f"{p}/doc-templates/{as_dotx['id']}/download").content
    Document(io.BytesIO(stored))  # opens: content type was converted
    assert [x["name"] for x in c.get(f"{p}/doc-templates")] == ["Corporate", "From dotx"]
    c.delete(f"{p}/doc-templates/{t['id']}", expect=403)
    m.delete(f"{p}/doc-templates/{t['id']}")
    assert len(c.get(f"{p}/doc-templates")) == 1


def test_export_into_a_customer_template(team, project, admin):
    c, m = team["contributor"], team["manager"]
    p = f"/projects/{project}"
    spec = new_item(c, project, "specification", "Sensor specification")
    h = new_item(c, project, "heading", "Interfaces", parent=spec["key"])
    sub = new_item(c, project, "heading", "Analog output", parent=h["key"])
    req = new_item(c, project, "system_requirement", "Provide 4-20 mA", parent=sub["key"], description=RICH)
    t = _upload(m, project, "ACME", _minimal_template())
    assert "unknown_thing" in t["unknown"]
    out = Document(io.BytesIO(c.get(f"{p}/items/{spec['key']}/export.docx?template={t['id']}").content))
    texts = _texts(out)
    # placeholders filled, also the split one; unknown placeholders are left visible for the template author
    assert texts[0].startswith("Sensor specification (" + spec["key"] + " A.1, In Work) by Carla")
    assert "{{unknown_thing}}" in texts[0]
    assert out.sections[0].header.paragraphs[0].text == "ACME Corp. — Test Product"
    body = "\n".join(texts)
    assert "{{title}}" not in body and "{{content}}" not in body and "{{toc}}" not in body
    # the content sits where the marker was — before the closing page — and starts after the table of contents
    i_req = next(i for i, x in enumerate(texts) if req["key"] in x)
    assert i_req < texts.index("END OF DOCUMENT — signatures on the next page")
    toc = [f for f in out.element.body.iter(qn("w:fldSimple")) if "TOC" in f.get(qn("w:instr"))]
    assert len(toc) == 1
    # styles the template lacked were copied in, list paragraphs keep working numbering
    names = {s.name for s in out.styles}
    assert {"Heading 2", "Heading 3", "List Bullet", "List Number", "Light Grid Accent 1"} <= names
    heading2 = next(par for par in out.paragraphs if par.text.endswith("Analog output"))
    assert heading2.style.name == "Heading 2"
    bullets = [par for par in out.paragraphs if par.text in ("first point", "second point")]
    assert bullets and all(par.style.name == "List Bullet" for par in bullets)
    num_ids = {n.get(qn("w:numId")) for n in out.part.numbering_part.element.findall(qn("w:num"))}
    for name in ("List Bullet", "List Number"):
        num = out.styles[name].element.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
        assert num is not None and num.get(qn("w:val")) in num_ids
    assert any(tb.cell(1, 1).text == "4-20 mA" for tb in out.tables)
    # the standard layout is unchanged without a template
    plain = Document(io.BytesIO(c.get(f"{p}/items/{spec['key']}/export.docx").content))
    assert plain.paragraphs[0].text == "Sensor specification" and "ACME" not in plain.sections[0].header.paragraphs[0].text
    # the round trip works inside a template too
    review = Document(io.BytesIO(c.get(f"{p}/items/{spec['key']}/export.docx?roundtrip=edit&template={t['id']}").content))
    tags = {s.find(f"{qn('w:sdtPr')}/{qn('w:tag')}").get(qn("w:val")) for s in review.element.body.iter(qn("w:sdt"))}
    assert any(tag.endswith(":desc") for tag in tags)


def test_templates_stay_in_their_project(team, project, admin):
    m = team["manager"]
    t = _upload(m, project, "ACME", _minimal_template())
    admin.post("/projects", {"key": "OTHER", "name": "Other", "template": "systems_engineering"})
    other = new_item(admin, "OTHER", "specification", "Other spec")
    admin.get(f"/projects/OTHER/items/{other['key']}/export.docx?template={t['id']}", expect=404)
