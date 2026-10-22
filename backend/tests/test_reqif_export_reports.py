import io
import json
import zipfile

from conftest import approve, new_item
from docx import Document

DOORS_REQIF = """<?xml version="1.0" encoding="UTF-8"?>
<REQ-IF xmlns="http://www.omg.org/spec/ReqIF/20110401/reqif.xsd" xmlns:xhtml="http://www.w3.org/1999/xhtml">
  <THE-HEADER><REQ-IF-HEADER IDENTIFIER="h1"><TITLE>Customer requirements</TITLE></REQ-IF-HEADER></THE-HEADER>
  <CORE-CONTENT><REQ-IF-CONTENT>
    <DATATYPES>
      <DATATYPE-DEFINITION-STRING IDENTIFIER="dt-s" LONG-NAME="String" MAX-LENGTH="255"/>
      <DATATYPE-DEFINITION-XHTML IDENTIFIER="dt-x" LONG-NAME="XHTML"/>
      <DATATYPE-DEFINITION-ENUMERATION IDENTIFIER="dt-prio" LONG-NAME="Priority">
        <SPECIFIED-VALUES>
          <ENUM-VALUE IDENTIFIER="ev-must" LONG-NAME="Must"><PROPERTIES><EMBEDDED-VALUE KEY="0" OTHER-CONTENT="1"/></PROPERTIES></ENUM-VALUE>
          <ENUM-VALUE IDENTIFIER="ev-should" LONG-NAME="Should"><PROPERTIES><EMBEDDED-VALUE KEY="1" OTHER-CONTENT="2"/></PROPERTIES></ENUM-VALUE>
        </SPECIFIED-VALUES>
      </DATATYPE-DEFINITION-ENUMERATION>
    </DATATYPES>
    <SPEC-TYPES>
      <SPEC-OBJECT-TYPE IDENTIFIER="t-req" LONG-NAME="Customer Requirement">
        <SPEC-ATTRIBUTES>
          <ATTRIBUTE-DEFINITION-STRING IDENTIFIER="a-id" LONG-NAME="ReqIF.ForeignID"><TYPE><DATATYPE-DEFINITION-STRING-REF>dt-s</DATATYPE-DEFINITION-STRING-REF></TYPE></ATTRIBUTE-DEFINITION-STRING>
          <ATTRIBUTE-DEFINITION-STRING IDENTIFIER="a-chap" LONG-NAME="ReqIF.ChapterName"><TYPE><DATATYPE-DEFINITION-STRING-REF>dt-s</DATATYPE-DEFINITION-STRING-REF></TYPE></ATTRIBUTE-DEFINITION-STRING>
          <ATTRIBUTE-DEFINITION-XHTML IDENTIFIER="a-text" LONG-NAME="ReqIF.Text"><TYPE><DATATYPE-DEFINITION-XHTML-REF>dt-x</DATATYPE-DEFINITION-XHTML-REF></TYPE></ATTRIBUTE-DEFINITION-XHTML>
          <ATTRIBUTE-DEFINITION-ENUMERATION IDENTIFIER="a-prio" LONG-NAME="Priority" MULTI-VALUED="false"><TYPE><DATATYPE-DEFINITION-ENUMERATION-REF>dt-prio</DATATYPE-DEFINITION-ENUMERATION-REF></TYPE></ATTRIBUTE-DEFINITION-ENUMERATION>
          <ATTRIBUTE-DEFINITION-STRING IDENTIFIER="a-oem" LONG-NAME="OEM Status"><TYPE><DATATYPE-DEFINITION-STRING-REF>dt-s</DATATYPE-DEFINITION-STRING-REF></TYPE></ATTRIBUTE-DEFINITION-STRING>
        </SPEC-ATTRIBUTES>
      </SPEC-OBJECT-TYPE>
      <SPEC-RELATION-TYPE IDENTIFIER="rt-ref" LONG-NAME="refines"/>
      <SPECIFICATION-TYPE IDENTIFIER="st" LONG-NAME="Spec"/>
    </SPEC-TYPES>
    <SPEC-OBJECTS>
      <SPEC-OBJECT IDENTIFIER="o-h1"><TYPE><SPEC-OBJECT-TYPE-REF>t-req</SPEC-OBJECT-TYPE-REF></TYPE><VALUES>
        <ATTRIBUTE-VALUE-STRING THE-VALUE="1 General"><DEFINITION><ATTRIBUTE-DEFINITION-STRING-REF>a-chap</ATTRIBUTE-DEFINITION-STRING-REF></DEFINITION></ATTRIBUTE-VALUE-STRING>
      </VALUES></SPEC-OBJECT>
      <SPEC-OBJECT IDENTIFIER="o-r1"><TYPE><SPEC-OBJECT-TYPE-REF>t-req</SPEC-OBJECT-TYPE-REF></TYPE><VALUES>
        <ATTRIBUTE-VALUE-STRING THE-VALUE="CR-0815"><DEFINITION><ATTRIBUTE-DEFINITION-STRING-REF>a-id</ATTRIBUTE-DEFINITION-STRING-REF></DEFINITION></ATTRIBUTE-VALUE-STRING>
        <ATTRIBUTE-VALUE-XHTML><DEFINITION><ATTRIBUTE-DEFINITION-XHTML-REF>a-text</ATTRIBUTE-DEFINITION-XHTML-REF></DEFINITION>
          <THE-VALUE><xhtml:div><xhtml:p>The sensor shall measure <xhtml:b>0 to 10 bar</xhtml:b>.</xhtml:p><xhtml:ul><xhtml:li>gauge</xhtml:li><xhtml:li>absolute</xhtml:li></xhtml:ul></xhtml:div></THE-VALUE>
        </ATTRIBUTE-VALUE-XHTML>
        <ATTRIBUTE-VALUE-ENUMERATION><DEFINITION><ATTRIBUTE-DEFINITION-ENUMERATION-REF>a-prio</ATTRIBUTE-DEFINITION-ENUMERATION-REF></DEFINITION><VALUES><ENUM-VALUE-REF>ev-must</ENUM-VALUE-REF></VALUES></ATTRIBUTE-VALUE-ENUMERATION>
        <ATTRIBUTE-VALUE-STRING THE-VALUE="accepted"><DEFINITION><ATTRIBUTE-DEFINITION-STRING-REF>a-oem</ATTRIBUTE-DEFINITION-STRING-REF></DEFINITION></ATTRIBUTE-VALUE-STRING>
      </VALUES></SPEC-OBJECT>
      <SPEC-OBJECT IDENTIFIER="o-r2"><TYPE><SPEC-OBJECT-TYPE-REF>t-req</SPEC-OBJECT-TYPE-REF></TYPE><VALUES>
        <ATTRIBUTE-VALUE-XHTML><DEFINITION><ATTRIBUTE-DEFINITION-XHTML-REF>a-text</ATTRIBUTE-DEFINITION-XHTML-REF></DEFINITION>
          <THE-VALUE><xhtml:div>Output 4-20 mA</xhtml:div></THE-VALUE></ATTRIBUTE-VALUE-XHTML>
        <ATTRIBUTE-VALUE-ENUMERATION><DEFINITION><ATTRIBUTE-DEFINITION-ENUMERATION-REF>a-prio</ATTRIBUTE-DEFINITION-ENUMERATION-REF></DEFINITION><VALUES><ENUM-VALUE-REF>ev-should</ENUM-VALUE-REF></VALUES></ATTRIBUTE-VALUE-ENUMERATION>
      </VALUES></SPEC-OBJECT>
    </SPEC-OBJECTS>
    <SPEC-RELATIONS>
      <SPEC-RELATION IDENTIFIER="r1"><TYPE><SPEC-RELATION-TYPE-REF>rt-ref</SPEC-RELATION-TYPE-REF></TYPE>
        <SOURCE><SPEC-OBJECT-REF>o-r2</SPEC-OBJECT-REF></SOURCE><TARGET><SPEC-OBJECT-REF>o-r1</SPEC-OBJECT-REF></TARGET></SPEC-RELATION>
    </SPEC-RELATIONS>
    <SPECIFICATIONS>
      <SPECIFICATION IDENTIFIER="s1" LONG-NAME="Customer specification"><TYPE><SPECIFICATION-TYPE-REF>st</SPECIFICATION-TYPE-REF></TYPE>
        <CHILDREN><SPEC-HIERARCHY IDENTIFIER="sh1"><OBJECT><SPEC-OBJECT-REF>o-h1</SPEC-OBJECT-REF></OBJECT>
          <CHILDREN>
            <SPEC-HIERARCHY IDENTIFIER="sh2"><OBJECT><SPEC-OBJECT-REF>o-r1</SPEC-OBJECT-REF></OBJECT></SPEC-HIERARCHY>
            <SPEC-HIERARCHY IDENTIFIER="sh3"><OBJECT><SPEC-OBJECT-REF>o-r2</SPEC-OBJECT-REF></OBJECT></SPEC-HIERARCHY>
          </CHILDREN></SPEC-HIERARCHY></CHILDREN>
      </SPECIFICATION>
    </SPECIFICATIONS>
  </REQ-IF-CONTENT></CORE-CONTENT>
</REQ-IF>"""


def _import(api, project, content: bytes, filename="customer.reqif", dry_run=True, mapping=None, default_type=None,
            expect=200):
    data = {"dry_run": str(dry_run).lower()}
    if mapping:
        data["mapping"] = json.dumps(mapping)
    if default_type:
        data["default_type"] = default_type
    return api.post(f"/projects/{project}/import-reqif", files={"file": (filename, io.BytesIO(content), "application/xml")},
                    data=data, expect=expect)


def test_reqif_import_doors_style(team, project):
    c = team["contributor"]
    p = f"/projects/{project}"
    dry = _import(c, project, DOORS_REQIF.encode())
    # Without a matching item type the objects are skipped until a type is chosen; the specification is created.
    assert dry["applied"] is False and dry["summary"]["skipped"] == 3 and dry["summary"]["created"] == 1
    t = dry["analysis"]["types"][0]
    assert t["name"] == "Customer Requirement" and t["mapped"] is None  # no matching type -> choose one
    assert dry["analysis"]["relations"][0]["mapped"] is None
    assert c.get(f"{p}/items")["total"] == 0
    mapping = {"types": {"Customer Requirement": "stakeholder_requirement"}, "relations": {"refines": "relates_to"}}
    dry = _import(c, project, DOORS_REQIF.encode(), mapping=mapping)
    attrs = dry["analysis"]["types"][0]["attributes"]
    assert attrs["fields"] == {"Priority": "priority"} and attrs["unmapped"] == ["OEM Status"]
    res = _import(c, project, DOORS_REQIF.encode(), dry_run=False, mapping=mapping)
    assert res["applied"] is True and res["summary"]["links"] == 1, res
    items = {i["title"]: i for i in c.get(f"{p}/items?type=stakeholder_requirement")["items"]}
    r1 = next(i for title, i in items.items() if title.startswith("The sensor shall measure"))
    detail = c.get(f"{p}/items/{r1['key']}")
    assert "<strong>0 to 10 bar</strong>" in detail["description"] and "<li><p>gauge</p></li>" in detail["description"]
    assert detail["fields"]["priority"] == "must" and "1 General" not in items  # imported as heading
    assert res["summary"]["headings"] == 1
    spec = c.get(f"{p}/items?type=specification")["items"][0]
    assert spec["title"] == "Customer specification"
    tree = c.get(f"{p}/items/{spec['key']}/tree")
    assert tree["children"][0]["title"] == "1 General" and len(tree["children"][0]["children"]) == 2
    refs = c.get(f"{p}/items/{r1['key']}/external-refs")
    assert refs[0]["system"] == "reqif" and refs[0]["external_id"] == "o-r1"

    # Re-import of a changed file updates the existing item instead of duplicating it.
    changed = DOORS_REQIF.replace("0 to 10 bar", "0 to 16 bar").encode()
    again = _import(c, project, changed, dry_run=False, mapping=mapping)
    assert again["summary"]["created"] == 0 and again["summary"]["updated"] == 1
    assert "16 bar" in c.get(f"{p}/items/{r1['key']}")["description"]
    assert c.get(f"{p}/items?type=stakeholder_requirement")["total"] == 2


def test_reqif_reimport_respects_frozen_items(team, project):
    c, r = team["contributor"], team["reviewer"]
    p = f"/projects/{project}"
    mapping = {"types": {"Customer Requirement": "system_requirement"}}
    _import(c, project, DOORS_REQIF.encode(), dry_run=False, mapping=mapping)
    item = next(i for i in c.get(f"{p}/items?type=system_requirement")["items"] if "sensor" in i["title"])
    c.post(f"{p}/items/{item['key']}/transitions/submit", {})
    approve(item["key"], project, c, r)
    res = _import(c, project, DOORS_REQIF.replace("0 to 10 bar", "0 to 25 bar").encode(), dry_run=False,
                  mapping=mapping)
    assert res["summary"]["frozen"] == 1 and any(item["key"] in w for w in res["warnings"])
    assert "10 bar" in c.get(f"{p}/items/{item['key']}")["description"]


def test_reqif_round_trip_between_projects(team, project, admin):
    c = team["contributor"]
    p = f"/projects/{project}"
    spec = new_item(c, project, "specification", "System spec")
    h = new_item(c, project, "heading", "1 Measurement", parent=spec["key"])
    a = new_item(c, project, "system_requirement", "Measure pressure", parent=h["key"],
                 description="Para one\n\nPara two", fields={"priority": "must", "safety_relevant": True})
    b = new_item(c, project, "system_requirement", "Output current", parent=h["key"])
    c.post(f"{p}/links", {"type": "derives_from", "source": b["key"], "target": a["key"]})
    exported = c.get(f"{p}/items/{spec['key']}/reqif").content

    # Into the same project: the file maps back onto the original items (no duplicates).
    same = _import(c, project, exported, dry_run=False)
    assert same["summary"]["created"] == 0 and same["summary"]["unchanged"] >= 3

    # Into another project: full structure, fields and links recreated.
    admin.post("/projects", {"key": "COPY", "name": "Copy", "template": "systems_engineering"})
    res = _import(admin, "COPY", exported, dry_run=False)
    assert res["applied"] is True and res["summary"]["links"] == 1, res
    copied = {i["title"]: i for i in admin.get("/projects/COPY/items?type=system_requirement")["items"]}
    ca = admin.get(f"/projects/COPY/items/{copied['Measure pressure']['key']}")
    assert ca["description"] == "<p>Para one</p><p>Para two</p>"
    assert ca["fields"]["priority"] == "must" and ca["fields"]["safety_relevant"] is True
    spec_copy = admin.get("/projects/COPY/items?type=specification")["items"][0]
    tree = admin.get(f"/projects/COPY/items/{spec_copy['key']}/tree")
    assert [n["title"] for n in tree["children"][0]["children"]] == ["Measure pressure", "Output current"]

    zipped = io.BytesIO()
    with zipfile.ZipFile(zipped, "w") as z:
        z.writestr("spec.reqif", exported)
    assert _import(admin, "COPY", zipped.getvalue(), filename="spec.reqifz")["summary"]["created"] == 0


def test_reqif_invalid_input(team, project):
    c = team["contributor"]
    _import(c, project, b"<not reqif", expect=422)
    _import(c, project, b"<root/>", expect=422)
    _import(c, project, b"PK\x03\x04garbage", filename="x.reqifz", expect=422)
    _import(c, project, DOORS_REQIF.encode(), mapping={"types": {"Customer Requirement": "nope"}}, expect=422)
    team["viewer"].post(f"/projects/{project}/import-reqif",
                        files={"file": ("a.reqif", io.BytesIO(DOORS_REQIF.encode()), "application/xml")}, expect=403)


def _document(c, project):
    spec = new_item(c, project, "specification", "System spec", description="Scope of the **system**.")
    h = new_item(c, project, "heading", "1 Measurement", parent=spec["key"])
    r1 = new_item(c, project, "system_requirement", "Measure pressure", parent=h["key"],
                  description="Range 0–10 bar, accuracy ±0.1 %\n\n- gauge\n- absolute\n\n| A | B |\n|---|---|\n| 1 | 2 |",
                  fields={"priority": "must"}, applicability="OUT_HART")
    stk = new_item(c, project, "stakeholder_requirement", "Monitor pressure")
    c.post(f"/projects/{project}/links", {"type": "derives_from", "source": r1["key"], "target": stk["key"]})
    return spec, h, r1


def test_word_export(team, project):
    c = team["contributor"]
    spec, h, r1 = _document(c, project)
    resp = c.get(f"/projects/{project}/items/{spec['key']}/export.docx")
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
    doc = Document(io.BytesIO(resp.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "System spec" in text and "1 Measurement" in text and r1["key"] in text
    assert "Range 0–10 bar" in text and "derives from" in text and "Applies to: OUT_HART" in text
    assert any(p.style.name.startswith("List") for p in doc.paragraphs)
    assert len(doc.tables) >= 3  # metadata, markdown table, fields
    lean = Document(io.BytesIO(c.get(f"/projects/{project}/items/{spec['key']}/export.docx?fields=false&links=false").content))
    assert "derives from" not in "\n".join(p.text for p in lean.paragraphs)


def test_pdf_export_live_and_baseline(team, project):
    c, r = team["contributor"], team["reviewer"]
    spec, h, r1 = _document(c, project)
    pdf = c.get(f"/projects/{project}/items/{spec['key']}/export.pdf")
    assert pdf.headers["content-type"] == "application/pdf" and pdf.content.startswith(b"%PDF")
    assert len(pdf.content) > 2000
    bl = r.post(f"/projects/{project}/baselines", {"name": "Draft 1", "scope": spec["key"]})
    c.patch(f"/projects/{project}/items/{r1['key']}", {"title": "Measure process pressure"})
    old = Document(io.BytesIO(c.get(f"/projects/{project}/items/{spec['key']}/export.docx?baseline={bl['id']}").content))
    old_text = "\n".join(p.text for p in old.paragraphs)
    assert "Measure pressure" in old_text and "Measure process pressure" not in old_text
    assert "Baseline: Draft 1" in "\n".join(cell.text for t in old.tables for row in t.rows for cell in row.cells)
    assert c.get(f"/projects/{project}/items/{spec['key']}/export.pdf?baseline={bl['id']}").content.startswith(b"%PDF")
    other = r.post(f"/projects/{project}/baselines", {"name": "Only req", "scope": r1["key"]})
    c.get(f"/projects/{project}/items/{spec['key']}/export.pdf?baseline={other['id']}", expect=404)


# --- reports ------------------------------------------------------------------------------------


def test_report_crud_and_visibility(team, project):
    c, m, v = team["contributor"], team["manager"], team["viewer"]
    p = f"/projects/{project}"
    assert {t["key"] for t in c.get(f"{p}/report-templates")} >= {"kpi_overview", "change_management"}
    rep = c.post(f"{p}/reports", {"title": "My KPIs", "template": "kpi_overview"}, expect=201)
    assert rep["widget_count"] == 8 and rep["mine"] is True
    private = c.post(f"{p}/reports", {"title": "Scratch", "visibility": "private", "widgets": [
        {"id": "w1", "type": "markdown", "text": "hello"}]}, expect=201)
    assert {r["title"] for r in v.get(f"{p}/reports")} == {"My KPIs"}
    v.get(f"{p}/reports/{private['id']}", expect=404)
    v.post(f"{p}/reports", {"title": "x"}, expect=403)
    team["reviewer"].put(f"{p}/reports/{rep['id']}", {"title": "Hijack"}, expect=403)
    m.put(f"{p}/reports/{rep['id']}", {"title": "Team KPIs", "widgets": rep["widgets"][:2]})
    assert c.get(f"{p}/reports/{rep['id']}")["widget_count"] == 2
    c.post(f"{p}/reports", {"title": "Dup", "widgets": [{"id": "a", "type": "markdown"}, {"id": "a", "type": "kpi"}]},
           expect=422)
    c.post(f"{p}/reports", {"title": "Bad", "widgets": [{"id": "a", "type": "pie-in-the-sky"}]}, expect=422)
    c.post(f"{p}/reports", {"title": "Bad", "template": "nope"}, expect=422)
    c.delete(f"{p}/reports/{private['id']}")


def test_widget_evaluation(team, project):
    c = team["contributor"]
    p = f"/projects/{project}"
    for title, prio in (("A", "must"), ("B", "must"), ("C", "should")):
        new_item(c, project, "system_requirement", title, fields={"priority": prio})
    tc = new_item(c, project, "test_case", "T")
    c.post(f"{p}/links", {"type": "verifies", "source": tc["key"], "target": "SYS-1"})

    def run(widget):
        return c.post(f"{p}/reports/preview-widget", {"id": "w", **widget})

    assert run({"type": "kpi", "filter": {"types": ["system_requirement"]}})["value"] == 3
    pct = run({"type": "kpi", "metric": "percent", "filter": {"types": ["system_requirement"]},
               "percent_of": {"fields": {"priority": "must"}}, "target": 50})
    assert pct["value"] == 67 and pct["numerator"] == 2 and pct["status"] == "good"
    chart = run({"type": "chart", "group_by": "field:priority", "stack_by": "state",
                 "filter": {"types": ["system_requirement"]}})
    assert {g["label"]: g["total"] for g in chart["groups"]} == {"Must": 2, "Should": 1}
    assert chart["stacks"] == [{"key": "draft", "label": "Draft"}]
    table = run({"type": "table", "filter": {"types": ["system_requirement"]}, "columns": ["key", "title", "field:priority"]})
    assert table["rows"][0] == {"key": "SYS-1", "title": "A", "field:priority": "Must"}
    trend = run({"type": "trend", "series": ["created", "released", "changed"], "cumulative": True,
                 "filter": {"types": ["system_requirement"]}})
    assert trend["series"]["created"][-1] == 3 and len(trend["periods"]) == 12
    cov = run({"type": "coverage", "coverage_type": "system_requirement"})
    assert cov["total"] == 3 and cov["with_tests_pct"] == 33
    mx = run({"type": "matrix", "rows": "system_requirement", "cols": "test_case"})
    assert mx["rows_covered_pct"] == 33
    assert run({"type": "markdown", "text": "**hi**"})["text"] == "**hi**"
    err = run({"type": "kpi", "filter": {"types": ["no_such_type"]}})
    assert err["ok"] is False and "no_such_type" in err["error"]
    assert run({"type": "chart", "group_by": "colour"})["ok"] is False


def test_kpis_reuse_lead_time_retest_ratio(team, project):
    c, m = team["contributor"], team["manager"]
    p = f"/projects/{project}"
    # Two products sharing one reused module; one new module used once.
    p1 = new_item(c, project, "system_element", "Product 1", fields={"level": "product"})
    p2 = new_item(c, project, "system_element", "Product 2", fields={"level": "product"})
    shared = new_item(c, project, "system_element", "Sensor cell", fields={"level": "module", "reuse_strategy": "reused"})
    new_mod = new_item(c, project, "system_element", "HART modem", applicability="HART",
                       fields={"level": "module", "reuse_strategy": "new"})
    for parent, child in ((p1, shared), (p2, shared), (p2, new_mod)):
        c.post(f"{p}/links", {"type": "contains", "source": parent["key"], "target": child["key"]})
    reuse = c.get(f"{p}/kpis/reuse")
    assert reuse["modules"] == 2 and reuse["reuse_rate_pct"] == 50 and reuse["multi_use_pct"] == 50
    assert reuse["rows"][0]["key"] == shared["key"] and reuse["rows"][0]["products"] == 2

    # Lead time: a change request approved by the CCB, and an implementing change order released.
    req = new_item(c, project, "system_requirement", "R")
    c.post(f"{p}/items/{req['key']}/transitions/submit", {})
    approve(req["key"], project, c, team["reviewer"])
    cr = new_item(c, project, "change_request", "Change R")
    c.post(f"{p}/links", {"type": "affects", "source": cr["key"], "target": req["key"]})
    c.post(f"{p}/items/{cr['key']}/transitions/analyse", {})
    approve(cr["key"], project, c, m)
    co = new_item(c, project, "change_order", "Implement")
    c.post(f"{p}/links", {"type": "implements", "source": co["key"], "target": cr["key"]})
    c.post(f"{p}/links", {"type": "affects", "source": co["key"], "target": req["key"]})
    c.post(f"{p}/items/{co['key']}/transitions/start_work", {})
    c.post(f"{p}/items/{req['key']}/revise", {"change_order": co["key"]})
    c.post(f"{p}/items/{co['key']}/transitions/submit", {})
    approve(co["key"], project, c, m, transition="release")
    lead = c.get(f"{p}/kpis/change-lead-time")
    assert lead["types"]["change_request"]["count"] == 1 and lead["types"]["change_request"]["open"] == 0
    assert lead["types"]["change_order"]["lead_days"]["avg"] is not None
    assert lead["end_to_end"]["avg"] is not None and len(lead["trend"]) == 1

    # Re-test effort and variant ratio.
    new_item(c, project, "feature", "HART", fields={"code": "HART"})
    new_item(c, project, "variant", "Basic", fields={"selection": []})
    new_item(c, project, "variant", "HART variant", fields={"selection": ["HART"]})
    t1 = new_item(c, project, "test_case", "Common test")
    new_item(c, project, "test_case", "HART test", applicability="HART")
    run_item = new_item(c, project, "test_run", "Run")
    recs = c.post(f"{p}/test-runs/{run_item['key']}/records", {"test_cases": [t1["key"]]})
    c.patch(f"{p}/test-runs/{run_item['key']}/records/{recs[0]['id']}", {"result": "passed", "duration_ms": 120000})
    retest = c.get(f"{p}/kpis/retest-effort")
    by_variant = {r["title"]: r for r in retest["variants"]}
    assert by_variant["Basic"]["applicable"] == 1 and by_variant["Basic"]["needs_retest"] == 0
    assert by_variant["HART variant"]["applicable"] == 2 and by_variant["HART variant"]["needs_retest"] == 1
    c.patch(f"{p}/items/{t1['key']}", {"title": "Common test v2"})  # newer test case version -> re-test
    assert {r["title"]: r for r in c.get(f"{p}/kpis/retest-effort")["variants"]}["Basic"]["needs_retest"] == 1
    ratio = c.get(f"{p}/kpis/variant-ratio")
    assert ratio["offered_variants"] == 2 and ratio["modules_total"] == 2 and ratio["modules_variant_specific"] == 1
    assert {r["title"]: r["modules"] for r in ratio["rows"]} == {"Basic": 1, "HART variant": 2}
    c.get(f"{p}/kpis/unknown", expect=404)


def test_report_data_and_excel(team, project):
    from openpyxl import load_workbook

    c = team["contributor"]
    p = f"/projects/{project}"
    new_item(c, project, "system_requirement", "A", fields={"priority": "must"})
    rep = c.post(f"{p}/reports", {"title": "Status", "template": "requirements_status"}, expect=201)
    data = c.get(f"{p}/reports/{rep['id']}/data")
    assert set(data) == {w["id"] for w in rep["widgets"]} and all(d["ok"] for d in data.values())
    wb = load_workbook(io.BytesIO(c.get(f"{p}/reports/{rep['id']}/export.xlsx").content))
    assert wb.sheetnames[0] == "Report" and len(wb.sheetnames) >= 4
